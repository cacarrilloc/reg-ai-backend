"""Generate a comprehensive compliance-update DOCX for the full policy library."""

from __future__ import annotations

import asyncio
import json
import re
from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO
from typing import Any
from uuid import uuid4

import structlog
from docx import Document
from fastapi import HTTPException, status
from sqlalchemy.orm import Session

from app.ai import ChatMessage, LLMProvider, optional_llm_provider
from app.core.config import Settings, get_settings
from app.core.observability import langfuse_workflow_trace, normalize_trace_tags
from app.models.core import GapAnalysis, Policy
from app.schemas.policies import ComplianceReportResponse
from app.services.gap_analysis_service import gap_analysis_ids_cleared_by_reupload
from app.services.llm_json import extract_json_string_field
from app.services.llm_json import response_json as parse_response_json
from app.services.llm_persona import (
    internal_policy_analysis_workflow,
    with_internal_policy_analysis_persona,
)
from app.services.policy_types import REGULATORY_REFERENCE_TYPE
from app.services.storage_service import create_download_url, storage_path, upload_file

logger = structlog.get_logger(__name__)

COMPLIANCE_REPORT_BUCKET = 'compliance-reports'
DOCX_CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
MAX_POLICY_EXCERPT = 3_000
MAX_GAPS_FOR_LLM = 40
COMPLIANCE_REPORT_MAX_TOKENS = 2_048
LLM_COMPLIANCE_REPORT_MAX_ATTEMPTS = 3
LLM_COMPLIANCE_REPORT_RETRY_DELAY_SECONDS = 2.0

COMPLIANCE_REPORT_SYSTEM_PROMPT = with_internal_policy_analysis_persona(
    """Prepare a comprehensive compliance remediation plan for a US financial-services firm.

After Phase 1A (deep understanding of each supplied Internal Policy) and Phase 1B (deep
understanding of the identified gaps and regulatory bases), produce a plan to bring the firm's
policy library to full regulatory alignment. Do not draft from gap titles alone.

Return only strict, valid JSON with exactly this shape (emit phase1_brief first):
{
  "phase1_brief": {
    "policy_understanding": "concise synthesis of the Internal Policies in scope",
    "companion_understanding": "concise synthesis of open gaps and regulatory bases",
    "covered_well": "areas already adequately controlled",
    "silent_or_thin": "highest-priority policy gaps remaining"
  },
  "executive_summary": "concise overview for the CCO",
  "required_changes": [
    {
      "policy_name": "policy title",
      "priority": "HIGH" | "MEDIUM",
      "change_description": "specific amendment needed",
      "regulatory_basis": "trigger source and citation",
      "implementation_steps": ["step 1", "step 2"]
    }
  ],
  "cross_cutting_themes": ["theme 1", "theme 2"]
}

Keep executive_summary under 250 words. Limit required_changes to the most material items.
Escape double quotes and newlines inside JSON string values as \\n. Do not use trailing commas or markdown fences.
Do not invent regulations or policy text. Base every change on the supplied gap analyses and policies."""
)


@dataclass(frozen=True)
class ComplianceSnapshot:
    status: str
    score: float
    policy_count: int
    compliant_policy_count: int
    policies_with_open_gaps: int
    open_gap_count: int
    high_gap_count: int
    medium_gap_count: int


def calculate_compliance_snapshot(
    policies: list[Policy],
    gaps: list[tuple[GapAnalysis, Policy]],
) -> ComplianceSnapshot:
    """Calculate the report's point-in-time policy-library compliance indicator."""
    impacted_policy_ids = {policy.id for _gap, policy in gaps}
    high_gap_count = sum(1 for gap, _policy in gaps if gap.impact_score == 'HIGH')
    medium_gap_count = sum(1 for gap, _policy in gaps if gap.impact_score == 'MEDIUM')
    policy_count = len(policies)
    compliant_policy_count = max(policy_count - len(impacted_policy_ids), 0)
    score = round((compliant_policy_count / policy_count) * 100, 1) if policy_count else 0.0
    if high_gap_count:
        compliance_status = 'ACTION REQUIRED'
    elif medium_gap_count:
        compliance_status = 'NEEDS ATTENTION'
    else:
        compliance_status = 'NO OPEN MATERIAL GAPS'
    return ComplianceSnapshot(
        status=compliance_status,
        score=score,
        policy_count=policy_count,
        compliant_policy_count=compliant_policy_count,
        policies_with_open_gaps=len(impacted_policy_ids),
        open_gap_count=len(gaps),
        high_gap_count=high_gap_count,
        medium_gap_count=medium_gap_count,
    )


async def generate_compliance_report(
    org_id: str,
    db: Session | None,
    client: LLMProvider | None = None,
    *,
    raise_on_failure: bool = True,
) -> ComplianceReportResponse | None:
    """Create a DOCX compliance report covering all active policies and open gaps."""
    async with langfuse_workflow_trace(
        'compliance-report',
        user_id=org_id,
        tags=normalize_trace_tags('compliance-report'),
        metadata={'org_id': org_id},
        input_data={'org_id': org_id},
    ):
        return await _generate_compliance_report_impl(
            org_id,
            db,
            client=client,
            raise_on_failure=raise_on_failure,
        )


async def _generate_compliance_report_impl(
    org_id: str,
    db: Session | None,
    client: LLMProvider | None = None,
    *,
    raise_on_failure: bool = True,
) -> ComplianceReportResponse | None:
    if db is None:
        if raise_on_failure:
            raise HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail='Compliance report generation requires a configured database',
            )
        logger.warning('compliance_report_skipped', reason='database_unavailable', org_id=org_id)
        return None

    settings = get_settings()
    policies = (
        db.query(Policy)
        .filter(
            Policy.org_id == org_id,
            Policy.status == 'ACTIVE',
            Policy.type != REGULATORY_REFERENCE_TYPE,
        )
        .order_by(Policy.name.asc())
        .all()
    )
    if not policies:
        if raise_on_failure:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail='No active policies found for compliance report generation',
            )
        logger.warning('compliance_report_skipped', reason='no_active_policies', org_id=org_id)
        return None

    gaps = (
        db.query(GapAnalysis, Policy)
        .join(Policy, Policy.id == GapAnalysis.policy_id)
        .filter(
            GapAnalysis.org_id == org_id,
            Policy.org_id == org_id,
            Policy.status == 'ACTIVE',
            Policy.type != REGULATORY_REFERENCE_TYPE,
            GapAnalysis.impact_score.in_(['HIGH', 'MEDIUM']),
        )
        .order_by(Policy.name.asc(), GapAnalysis.created_at.desc())
        .all()
    )
    cleared_gap_ids = gap_analysis_ids_cleared_by_reupload(db, org_id)
    gaps = [
        (gap, policy)
        for gap, policy in gaps
        if gap.id not in cleared_gap_ids
    ]

    snapshot = calculate_compliance_snapshot(policies, gaps)
    report_data = await draft_compliance_plan(
        policies,
        gaps[:MAX_GAPS_FOR_LLM],
        settings,
        client,
        org_id=org_id,
    )
    report_id = str(uuid4())
    docx_bytes = build_compliance_report_docx(
        org_id,
        policies,
        gaps,
        report_data,
        snapshot=snapshot,
    )

    try:
        stored = upload_file(
            COMPLIANCE_REPORT_BUCKET,
            storage_path(org_id, 'compliance-reports', report_id, f'compliance_report_{report_id}.docx'),
            docx_bytes,
            DOCX_CONTENT_TYPE,
        )
        expires_in = 900
        download_url = create_download_url(stored.uri, expires_in)
    except HTTPException as exc:
        logger.warning(
            'compliance_report_storage_failed',
            org_id=org_id,
            status_code=exc.status_code,
            detail=str(exc.detail),
        )
        if raise_on_failure:
            raise
        return None
    except Exception as exc:
        logger.warning(
            'compliance_report_storage_failed',
            org_id=org_id,
            error_type=type(exc).__name__,
            error=str(exc)[:500],
        )
        if raise_on_failure:
            raise HTTPException(
                status_code=status.HTTP_502_BAD_GATEWAY,
                detail='Unable to store compliance report',
            ) from exc
        return None

    return ComplianceReportResponse(
        report_id=report_id,
        org_id=org_id,
        file_url=stored.uri,
        download_url=download_url,
        expires_in=expires_in,
        policy_count=len(policies),
        gap_count=len(gaps),
        compliance_status=snapshot.status,
        compliance_score=snapshot.score,
        generated_at=datetime.now(UTC),
    )


async def draft_compliance_plan(
    policies: list[Policy],
    gaps: list[tuple[GapAnalysis, Policy]],
    settings: Settings,
    client: LLMProvider | None,
    *,
    org_id: str | None = None,
) -> dict[str, Any]:
    if client is None:
        client = optional_llm_provider(org_id, 'analysis', platform_model=settings.anthropic_opus_model)
    if client is None:
        return fallback_compliance_plan(policies, gaps)

    prompt = json.dumps(
        {
            'analysis_workflow': internal_policy_analysis_workflow(),
            'active_policies': [
                {'name': policy.name, 'version': policy.version, 'type': policy.type}
                for policy in policies
            ],
            'identified_gaps': [
                {
                    'policy_name': policy.name,
                    'policy_version': policy.version,
                    'impact_score': gap.impact_score,
                    'trigger_type': gap.trigger_type,
                    'gap_summary': gap.gap_summary,
                    'citations': gap.regulatory_citations,
                }
                for gap, policy in gaps
            ],
        },
        ensure_ascii=True,
    )
    last_error: Exception | None = None
    for attempt in range(1, LLM_COMPLIANCE_REPORT_MAX_ATTEMPTS + 1):
        try:
            response = await client.generate(
                [ChatMessage('user', prompt)],
                max_tokens=COMPLIANCE_REPORT_MAX_TOKENS,
                system=COMPLIANCE_REPORT_SYSTEM_PROMPT,
                temperature=0,
                cache_system=True,
                trace_name='compliance_report.llm',
            )
            return parse_response_json(
                response,
                loose_parser=parse_compliance_fields_loosely,
                recovery_log_event='compliance_report_json_recovered_with_loose_parser',
            )
        except Exception as exc:
            last_error = exc
            logger.warning(
                'compliance_report_llm_failed',
                error_type=type(exc).__name__,
                error=str(exc)[:500],
                attempt=attempt,
            )
            if attempt < LLM_COMPLIANCE_REPORT_MAX_ATTEMPTS:
                await asyncio.sleep(LLM_COMPLIANCE_REPORT_RETRY_DELAY_SECONDS * attempt)

    if settings.is_production:
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail='AI compliance report drafting failed',
        ) from last_error
    return fallback_compliance_plan(policies, gaps)


def parse_compliance_fields_loosely(text: str) -> dict[str, Any] | None:
    """Recover compliance report fields when Opus returns malformed or truncated JSON."""
    executive_summary = extract_json_string_field(text, ('executive_summary',))
    if not executive_summary:
        return None

    payload: dict[str, Any] = {
        'executive_summary': executive_summary,
        'required_changes': [],
        'cross_cutting_themes': [],
    }

    themes_match = re.search(r'"cross_cutting_themes"\s*:\s*\[(.*?)\]\s*\}', text, re.DOTALL)
    if themes_match is not None:
        payload['cross_cutting_themes'] = re.findall(r'"([^"\\]*(?:\\.[^"\\]*)*)"', themes_match.group(1))

    change_blocks = re.findall(
        r'\{\s*"policy_name"\s*:\s*"([^"]+)"\s*,\s*"priority"\s*:\s*"([^"]+)"\s*,'
        r'\s*"change_description"\s*:\s*"((?:[^"\\]|\\.)*)"',
        text,
        re.DOTALL,
    )
    for policy_name, priority, change_description in change_blocks:
        payload['required_changes'].append(
            {
                'policy_name': policy_name,
                'priority': priority,
                'change_description': (
                    change_description.replace('\\n', '\n').replace('\\"', '"').strip()
                ),
                'regulatory_basis': '',
                'implementation_steps': [
                    'Review the gap analysis finding',
                    'Generate and approve a tracked-change redline',
                    'Publish the updated policy version',
                ],
            }
        )
    return payload


def fallback_compliance_plan(
    policies: list[Policy],
    gaps: list[tuple[GapAnalysis, Policy]],
) -> dict[str, Any]:
    return {
        'executive_summary': (
            f'Reg AI identified {len(gaps)} material gap(s) across {len(policies)} active '
            'policies. A CCO must review each finding and approve specific amendments.'
        ),
        'required_changes': [
            {
                'policy_name': policy.name,
                'priority': gap.impact_score,
                'change_description': gap.gap_summary,
                'regulatory_basis': str(gap.regulatory_citations),
                'implementation_steps': [
                    'Review the gap analysis finding',
                    'Generate and approve a tracked-change redline',
                    'Publish the updated policy version',
                ],
            }
            for gap, policy in gaps
        ],
        'cross_cutting_themes': ['Supervision controls', 'Documentation and evidence'],
    }


def build_compliance_report_docx(
    org_id: str,
    policies: list[Policy],
    gaps: list[tuple[GapAnalysis, Policy]],
    report_data: dict[str, Any],
    *,
    snapshot: ComplianceSnapshot | None = None,
) -> bytes:
    snapshot = snapshot or calculate_compliance_snapshot(policies, gaps)
    document = Document()
    document.add_heading('Policy Library Compliance Update Report', level=0)
    document.add_paragraph(f'Organization: {org_id}')
    document.add_paragraph(f'Generated: {datetime.now(UTC).strftime("%Y-%m-%d %H:%M UTC")}')
    document.add_paragraph(
        'This report summarizes all changes required to bring the policy library to full '
        'regulatory compliance based on current gap analyses. All amendments require CCO review.'
    )

    document.add_heading('Current Compliance Status', level=1)
    document.add_paragraph(f'Status: {snapshot.status}')
    document.add_paragraph(f'{snapshot.score:.1f}% Compliance')
    document.add_paragraph(
        f'{snapshot.compliant_policy_count} of {snapshot.policy_count} active internal policies '
        'have no open HIGH or MEDIUM gap.'
    )
    document.add_paragraph(
        f'Open material gaps: {snapshot.open_gap_count} '
        f'({snapshot.high_gap_count} HIGH, {snapshot.medium_gap_count} MEDIUM) across '
        f'{snapshot.policies_with_open_gaps} policy/policies.'
    )
    document.add_paragraph(
        'Methodology: score = active internal policies with no open HIGH or MEDIUM gap '
        'divided by all active internal policies. A gap remains open until its approved '
        'revision is uploaded as the current policy version. This operational indicator '
        'is not a legal conclusion.'
    )

    document.add_heading('Executive Summary', level=1)
    document.add_paragraph(str(report_data.get('executive_summary') or ''))

    themes = report_data.get('cross_cutting_themes')
    if isinstance(themes, list) and themes:
        document.add_heading('Cross-Cutting Themes', level=1)
        for theme in themes:
            document.add_paragraph(str(theme), style='List Bullet')

    document.add_heading('Active Policy Inventory', level=1)
    for policy in policies:
        document.add_paragraph(
            f'{policy.name} (version {policy.version}, {policy.type})',
            style='List Bullet',
        )

    document.add_heading('Prioritized Remediation Plan', level=1)
    changes = report_data.get('required_changes')
    if isinstance(changes, list) and changes:
        for index, change in enumerate(changes, start=1):
            if not isinstance(change, dict):
                continue
            document.add_heading(
                f'{index}. {change.get("policy_name", "Policy")} — {change.get("priority", "MEDIUM")}',
                level=2,
            )
            document.add_paragraph(str(change.get('change_description') or ''))
            basis = change.get('regulatory_basis')
            if basis:
                document.add_paragraph(f'Regulatory basis: {basis}')
            steps = change.get('implementation_steps')
            if isinstance(steps, list):
                document.add_paragraph('Implementation steps:')
                for step in steps:
                    document.add_paragraph(str(step), style='List Number')
    elif gaps:
        for index, (gap, policy) in enumerate(gaps, start=1):
            document.add_heading(
                f'{index}. {policy.name} — {gap.impact_score}',
                level=2,
            )
            document.add_paragraph(gap.gap_summary)
    else:
        document.add_paragraph(
            'No material gaps were identified. Continue monitoring regulatory publications.'
        )

    document.add_heading('Complete Open Gap Register', level=1)
    document.add_paragraph(
        'The following register is generated directly from the current database and includes '
        'every open HIGH and MEDIUM gap for active internal policies; it is not shortened by AI.'
    )
    if gaps:
        for index, (gap, policy) in enumerate(gaps, start=1):
            document.add_heading(
                f'{index}. {policy.name} (version {policy.version}) — {gap.impact_score}',
                level=2,
            )
            document.add_paragraph(f'Gap ID: {gap.id}')
            document.add_paragraph(
                f'Identified: {gap.created_at.strftime("%Y-%m-%d %H:%M UTC")}'
                if gap.created_at
                else 'Identified: date unavailable'
            )
            document.add_paragraph(
                f'Trigger: {gap.trigger_type} ({gap.trigger_id})'
            )
            if gap.rule:
                document.add_paragraph(f'Rule: {gap.rule}')
            document.add_paragraph(f'Gap: {gap.gap_summary}')
            if gap.fix:
                document.add_paragraph(f'Recommended fix: {gap.fix}')
            citations = gap.regulatory_citations
            if isinstance(citations, list) and citations:
                document.add_paragraph('Regulatory citations:')
                for citation in citations:
                    if isinstance(citation, dict):
                        citation_text = ' — '.join(
                            str(value)
                            for value in (
                                citation.get('source'),
                                citation.get('reference'),
                                citation.get('url'),
                            )
                            if value
                        )
                    else:
                        citation_text = str(citation)
                    if citation_text:
                        document.add_paragraph(citation_text, style='List Bullet')
    else:
        document.add_paragraph('No open HIGH or MEDIUM gaps were identified.')

    document.add_heading('CCO Review Checklist', level=1)
    checklist = [
        'Review each required change against the cited regulatory trigger',
        'Generate tracked-change redlines for affected DOCX policies',
        'Approve or reject each redline through the review queue',
        'Publish approved policy versions to the canonical library',
    ]
    for item in checklist:
        document.add_paragraph(item, style='List Number')

    output = BytesIO()
    document.save(output)
    return output.getvalue()
