from __future__ import annotations

import re
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from types import EllipsisType
from typing import Any, Literal

import structlog
from fastapi import HTTPException, UploadFile, status

from app.ai import AIError, optional_embedding_provider
from app.core.config import Settings
from app.db.session import get_tenant_org_id

DocumentType = Literal['pdf', 'docx']

PDF_MIME_TYPE = 'application/pdf'
DOCX_MIME_TYPE = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
LEGACY_DOC_MIME_TYPE = 'application/msword'
DEFAULT_CHUNK_SIZE = 3600
DEFAULT_CHUNK_OVERLAP = 400

logger = structlog.get_logger(__name__)


@dataclass(frozen=True)
class DocumentSection:
    """A named section of a parsed document."""

    label: str
    index: int
    text: str


@dataclass(frozen=True)
class ParsedPolicyDocument:
    """A parsed policy document with extracted text and sections."""

    filename: str
    document_type: DocumentType
    text: str
    sections: list[DocumentSection]


@dataclass(frozen=True)
class TextChunk:
    """A chunk of text extracted from a document with metadata."""

    index: int
    content: str
    metadata: dict[str, Any]


@dataclass(frozen=True)
class ChunkUnit:
    """Represents a unit of text belonging to a specific document section.

    Attributes:
        section_index: The index of the section this chunk belongs to.
        text: The text content of the chunk.
    """

    section_index: int
    text: str


async def read_upload_bytes(file: UploadFile, settings: Settings) -> bytes:
    """Read uploaded file bytes and validate size and non-emptiness.

    Args:
        file: The uploaded file to read.
        settings: Application settings containing upload limits.

    Returns:
        The raw bytes of the uploaded file.

    Raises:
        HTTPException: If the file is empty or exceeds the configured size limit.
    """
    contents = await file.read()
    if not contents:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST, detail='Policy document is empty')

    max_bytes = settings.max_policy_upload_mb * 1024 * 1024
    if len(contents) > max_bytes:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail=f'Policy document exceeds the {settings.max_policy_upload_mb} MB upload limit',
        )

    return contents


def parse_policy_document(filename: str, content_type: str | None, contents: bytes) -> ParsedPolicyDocument:
    document_type = detect_document_type(filename, content_type)
    sections = extract_pdf_sections(
        contents) if document_type == 'pdf' else extract_docx_sections(contents)
    text = normalize_text('\n\n'.join(section.text for section in sections))

    if not text:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail='No readable text was found in the policy document',
        )

    return ParsedPolicyDocument(filename=filename, document_type=document_type, text=text, sections=sections)


def detect_document_type(filename: str, content_type: str | None) -> DocumentType:
    suffix = Path(filename).suffix.lower()
    mime_type = (content_type or '').split(';', maxsplit=1)[0].strip().lower()

    if suffix == '.pdf' or mime_type == PDF_MIME_TYPE:
        return 'pdf'

    if suffix == '.docx' or mime_type == DOCX_MIME_TYPE:
        return 'docx'

    if suffix == '.doc' or mime_type == LEGACY_DOC_MIME_TYPE:
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail='Legacy .doc files are not supported yet. Upload a PDF or .docx policy document.',
        )

    raise HTTPException(
        status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
        detail='Unsupported policy document type. Upload a PDF or .docx file.',
    )


def canonical_document_content_type(document_type: DocumentType) -> str:
    if document_type == 'pdf':
        return PDF_MIME_TYPE
    return DOCX_MIME_TYPE


def extract_pdf_sections(contents: bytes) -> list[DocumentSection]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail='PDF parsing dependency is not installed',
        ) from exc

    try:
        reader = PdfReader(BytesIO(contents))
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail='Unable to read PDF policy document',
        ) from exc

    sections: list[DocumentSection] = []
    for page_number, page in enumerate(reader.pages, start=1):
        page_text = normalize_text(page.extract_text() or '')
        if page_text:
            sections.append(DocumentSection(
                label=f'page {page_number}', index=page_number, text=page_text))

    return sections


def extract_docx_sections(contents: bytes) -> list[DocumentSection]:
    try:
        from docx import Document
    except ImportError as exc:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail='Word document parsing dependency is not installed',
        ) from exc

    try:
        document = Document(BytesIO(contents))
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail='Unable to read Word policy document',
        ) from exc

    parts: list[str] = []
    for paragraph in document.paragraphs:
        paragraph_text = normalize_text(paragraph.text)
        if paragraph_text:
            parts.append(paragraph_text)

    for table in document.tables:
        for row in table.rows:
            cells = [normalize_text(cell.text) for cell in row.cells]
            row_text = ' | '.join(cell for cell in cells if cell)
            if row_text:
                parts.append(row_text)

    text = normalize_text('\n\n'.join(parts))
    return [DocumentSection(label='document', index=1, text=text)] if text else []


def chunk_policy_document(
    parsed: ParsedPolicyDocument,
    chunk_size: int = DEFAULT_CHUNK_SIZE,
    chunk_overlap: int = DEFAULT_CHUNK_OVERLAP,
) -> list[TextChunk]:
    if chunk_size <= chunk_overlap:
        raise ValueError('chunk_size must be greater than chunk_overlap')

    section_labels = {
        section.index: section.label for section in parsed.sections}
    units = chunk_units(parsed.sections, chunk_size, chunk_overlap)
    chunks: list[TextChunk] = []
    current_parts: list[str] = []
    current_sections: list[int] = []

    for unit in units:
        if current_parts and joined_length([*current_parts, unit.text]) > chunk_size:
            previous_sections = current_sections.copy()
            append_chunk(chunks, parsed, current_parts,
                         current_sections, section_labels)
            overlap_text = chunks[-1].content[-chunk_overlap:
                                              ].strip() if chunk_overlap else ''
            if overlap_text and joined_length([overlap_text, unit.text]) <= chunk_size:
                current_parts = [overlap_text]
                current_sections = [previous_sections[-1]
                                    ] if previous_sections else []
            else:
                current_parts = []
                current_sections = []

        current_parts.append(unit.text)
        current_sections.append(unit.section_index)

    append_chunk(chunks, parsed, current_parts,
                 current_sections, section_labels)
    return chunks


def chunk_units(
    sections: list[DocumentSection],
    chunk_size: int,
    chunk_overlap: int,
) -> list[ChunkUnit]:
    units: list[ChunkUnit] = []
    for section in sections:
        paragraphs = [
            paragraph.strip()
            for paragraph in re.split(r'\n\s*\n', normalize_text(section.text))
            if paragraph.strip()
        ]
        for paragraph in paragraphs:
            units.extend(split_long_unit(
                section.index, paragraph, chunk_size, chunk_overlap))
    return units


def split_long_unit(section_index: int, text: str, chunk_size: int, chunk_overlap: int) -> list[ChunkUnit]:
    if len(text) <= chunk_size:
        return [ChunkUnit(section_index=section_index, text=text)]

    units: list[ChunkUnit] = []
    step = chunk_size - chunk_overlap
    start = 0
    while start < len(text):
        piece = text[start: start + chunk_size].strip()
        if piece:
            units.append(ChunkUnit(section_index=section_index, text=piece))
        if start + chunk_size >= len(text):
            break
        start += step

    return units


def append_chunk(
    chunks: list[TextChunk],
    parsed: ParsedPolicyDocument,
    parts: list[str],
    section_indices: list[int],
    section_labels: dict[int, str],
) -> None:
    content = normalize_text('\n\n'.join(parts))
    if not content:
        return

    unique_sections = sorted(set(section_indices))
    labels = [section_labels[index]
              for index in unique_sections if index in section_labels]
    chunks.append(
        TextChunk(
            index=len(chunks),
            content=content,
            metadata={
                'document_type': parsed.document_type,
                'filename': parsed.filename,
                'section_start': unique_sections[0] if unique_sections else None,
                'section_end': unique_sections[-1] if unique_sections else None,
                'section_labels': labels,
                'char_count': len(content),
                'word_count': len(content.split()),
            },
        )
    )


def joined_length(parts: list[str]) -> int:
    if not parts:
        return 0
    return sum(len(part) for part in parts) + (2 * (len(parts) - 1))


def normalize_text(text: str) -> str:
    text = text.replace('\x00', '').replace('\r\n', '\n').replace('\r', '\n')
    normalized_lines = [re.sub(r'[ \t\f\v]+', ' ', line).strip()
                        for line in text.split('\n')]

    output: list[str] = []
    for line in normalized_lines:
        if line:
            output.append(line)
        elif output and output[-1] != '':
            output.append('')

    return '\n'.join(output).strip()


def chunk_publication_text(title: str, raw_text: str | None) -> list[TextChunk]:
    """Chunk a regulatory publication title and body for pgvector embedding."""
    text = normalize_text(f'{title}\n\n{raw_text or ""}')
    if not text:
        return []
    parsed = ParsedPolicyDocument(
        filename='publication.txt',
        document_type='docx',
        text=text,
        sections=[DocumentSection(label='publication', index=1, text=text)],
    )
    return chunk_policy_document(parsed)


async def generate_embeddings(
    chunks: list[TextChunk],
    settings: Settings,
    *,
    org_id: str | None | EllipsisType = ...,
) -> list[list[float]]:
    """Embed chunks with the org's embedding model (platform OpenAI by default).

    ``org_id`` defaults to the tenant in context; pass ``None`` for shared global content.
    """
    if not chunks:
        return []

    provider = optional_embedding_provider(get_tenant_org_id() if org_id is ... else org_id)
    if provider is None:
        if settings.is_production:
            raise HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail='An embedding model is required for policy document ingestion',
            )
        return []

    try:
        result = await provider.embed([chunk.content for chunk in chunks], trace_name='embeddings.generate')
    except AIError as exc:
        logger.warning('embedding_generation_failed', error_type=type(exc).__name__, model=provider.model_label)
        if not settings.is_production:
            return []
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail='Embedding request failed during policy ingestion',
        ) from exc
    logger.info('embedding_generation_completed', model=provider.model_label, total_embeddings=len(result.vectors))
    return result.vectors
