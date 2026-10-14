"""Per-format CCO deliverable. Does not use or modify the R1 redline pipeline."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from app.services.marketing_types import PERFORMANCE_DISCLAIMER, REPORT_HEADER


def output_kind_for_format(format_name: str, findings: list[dict[str, Any]]) -> str:
    recommended = {str(item.get('recommended_option') or '') for item in findings}
    if 'pull' in recommended and format_name == 'social':
        return 'pull_draft'
    if format_name in {'webpage', 'image', 'social'}:
        return 'change_list'
    if format_name == 'presentation':
        return 'tracked_docx'
    return 'tracked_docx'


def build_change_list_artifact(
    *,
    title: str,
    format_name: str,
    findings: list[dict[str, Any]],
    change_list: list[dict[str, Any]],
    rule_version: str,
) -> dict[str, Any]:
    return {
        'title': title,
        'format': format_name,
        'header': REPORT_HEADER,
        'performance_disclaimer': PERFORMANCE_DISCLAIMER,
        'rule_version': rule_version,
        'findings': findings,
        'change_list': change_list,
    }


def build_tracked_docx_bytes(
    *,
    title: str,
    body_text: str,
    findings: list[dict[str, Any]],
    change_list: list[dict[str, Any]],
    rule_version: str,
) -> bytes:
    from docx import Document
    from docx.shared import RGBColor

    document = Document()
    document.add_heading(title or 'Marketing review draft', level=1)
    draft_banner = document.add_paragraph('DRAFT — CCO REVIEW REQUIRED')
    draft_banner.runs[0].bold = True
    header = document.add_paragraph(REPORT_HEADER)
    header.runs[0].italic = True
    document.add_paragraph(PERFORMANCE_DISCLAIMER)
    document.add_paragraph(f'Rule store version: {rule_version}')
    if any(item.get('slide') for item in findings):
        document.add_paragraph(
            'Global redline of the full deck. Each slide was reviewed for text and diagrams '
            'against SEC Rule 206(4)-1, SEC Rule 204-2, FINRA Rule 2210, and FINRA Rules 2211–2216. '
            'Figures are presentation-only — the agent has not verified whether any number is accurate.'
        )
    document.add_heading('Working copy', level=2)
    document.add_paragraph(body_text or '(No extractable body text — see change list.)')
    document.add_heading('Keyed changes', level=2)
    for item in change_list:
        para = document.add_paragraph()
        location_prefix = ''
        if item.get('slide'):
            location_prefix = f"Slide {item.get('slide')} · "
        elif item.get('part'):
            location_prefix = f"C{item.get('marker')} · {item.get('part')} · "
        run = para.add_run(
            f"{location_prefix}[C{item.get('marker')}] {item.get('action')}: "
            f"{item.get('draft_text') or item.get('excerpt')}"
        )
        if item.get('action') == 'redline':
            run.font.color.rgb = RGBColor(0x9B, 0x1B, 0x30)
        _add_comment_like_note(para, str(item.get('rule') or ''))
    document.add_heading('Findings', level=2)
    for finding in findings:
        location_prefix = ''
        if finding.get('slide'):
            location_prefix = f"Slide {finding.get('slide')} — "
        elif finding.get('part'):
            location_prefix = f"C{finding.get('marker')} · {finding.get('part')} — "
        document.add_paragraph(
            f"{location_prefix}{finding.get('title') or finding.get('excerpt')} — "
            f"{finding.get('rule')}: {finding.get('why')} "
            f"[{finding.get('recommended_option')}]"
        )
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _add_comment_like_note(paragraph: Any, text: str) -> None:
    if not text:
        return
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragraph.add_run(f'  ({text})')
    run.italic = True
    rpr = run._r.get_or_add_rPr()
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '666666')
    rpr.append(color)
