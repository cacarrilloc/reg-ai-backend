"""Marketing content redlines for the CCO Review Queue.

Reuses the existing PolicyRevision / Review Queue path and the tracked-change
RULE / GAP / FIX comment helpers. Does not modify the internal-policy LLM pipeline.
"""

from __future__ import annotations

import re
from io import BytesIO
from typing import Any
from uuid import uuid4
from zipfile import ZipFile

import structlog
from fastapi import HTTPException, status
from sqlalchemy.orm import Session

from app.core.time import utc_now
from app.models.core import (
    GapAnalysis,
    MarketingContent,
    MarketingReview,
    Organization,
    Policy,
    PolicyRevision,
)
from app.schemas.review import PolicyRevisionResponse
from app.services.gap_analysis_service import format_gap_rule_fix_comment
from app.services.marketing_gap_bridge import (
    MARKETING_GAP_TRIGGER,
    persist_marketing_gap_analysis,
    selected_or_recommended_option,
)
from app.services.redline_document import normalize_space
from app.services.redline_generation_service import (
    DOCX_CONTENT_TYPE,
    REDLINE_BUCKET,
    persist_revision_redline_url,
    policy_revision_response,
    redline_filename,
)
from app.services.storage_service import storage_path, upload_file
from app.services.tracked_redline_document import (
    TrackedReplacement,
    apply_tracked_redline_overlay,
)

logger = structlog.get_logger(__name__)

PULL_REPLACEMENT = '[Pulled — do not publish this copy.]'
EDIT_OPTIONS = frozenset({'redline', 'add_disclosure', 'pull'})
APPEND_REDLINE_IDS = frozenset({
    'performance_gross_hero',
    'performance_periods',
    'performance_fee_basis',
})
SKIP_LAYOUT_REGIONS: frozenset[str] = frozenset()
HOMEPAGE_TITLE_SKIP = frozenset({'page title'})
FINDING_REGION_IDS = {
    'performance_gross_hero': 'hero',
    'performance_fee_basis': 'hero',
    'performance_periods': 'hero',
    'testimonial': 'testimonial',
    'third_party_rating': 'badge_row',
    'safety_claim': 'mid_page',
    'layout_buried': 'footer',
    'eliminate_claim': 'para_6',
    'client_outcome_2022': 'para_7',
    'educational_half': 'educational',
    'missing_disclosures': 'disclosures',
    'free_cta': 'closing',
    'cherry_pick': 'bar_chart',
    'gross_net_layout': 'slide_footer',
    'hypothetical_backtest': 'callout',
    'benchmark_comparison': 'slide_footer',
    'puffery_title': 'slide_title',
    'unsubstantiated_exclusivity': 'slide_copy',
    'unsubstantiated_statistic': 'slide_copy',
    'third_party_outcome': 'slide_copy',
}
PLACEHOLDER_EXCERPTS = frozenset({
    'testimonial carousel',
    'disclosures exist but are buried',
    'performance figure',
    'third-party ratings',
    'third party ratings',
})
EMPTY_DRAFTS = frozenset({
    '',
    'redline',
    'add disclosure',
    'add_disclosure',
    'pull',
    'held',
    'do_nothing',
    'do nothing',
})
LAYOUT_REGION_LABELS = {
    'hero': 'HERO',
    'badge_row': 'BADGE ROW',
    'badge': 'BADGE ROW',
    'testimonial': 'TESTIMONIAL CAROUSEL',
    'mid_page': 'MID-PAGE',
    'mid': 'MID-PAGE',
    'footer': 'FOOTER',
    'title': 'TITLE',
    'educational': 'PARAS 1–5, EDUCATIONAL',
    'para_6': 'PARA 6',
    'para_7': 'PARA 7',
    'closing': 'CLOSING',
    'button': 'BUTTON',
    'disclosures': 'NO DISCLOSURES',
    'slide_title': 'SLIDE TITLE',
    'bar_chart': 'BAR CHART',
    'callout': 'CALLOUT BOX',
    'slide_footer': 'BOTTOM RIGHT',
    'appendix': 'APPENDIX',
    'diagram': 'DIAGRAM',
    'gross_net': 'BAR CHART',
    'benchmark': 'BAR CHART',
    'slide_copy': 'SLIDE COPY',
}
VISUAL_PASS_FONT_NOTES = {
    'hero': 'full-width banner, 48pt headline',
    'badge_row': 'three large gold seals, directly under hero',
    'badge': 'three large gold seals, directly under hero',
    'testimonial': 'rotating, 24pt type',
    'mid_page': '16pt',
    'mid': '16pt',
    'footer': '7pt grey text on white, below the fold, after the sitemap',
    'title': 'Insights blog headline',
    'educational': 'paras 1–5 — educational, no firm mention',
    'para_6': 'promotional turn — offers advisory services',
    'para_7': '2022 client-outcome / 60/40 comparison',
    'closing': 'closing CTA',
    'button': 'Book my call button',
    'disclosures': 'no disclosures anywhere on the page',
    'slide_title': '32pt slide title',
    'bar_chart': 'dominant, left two-thirds of slide; two bars, one year',
    'callout': '20pt gold callout',
    'slide_footer': '11pt bottom right',
    'appendix': 'slide 14, 9pt',
    'diagram': 'embedded chart / image converted for vector comparison',
}
INGEST_FONT_NOTE_ALIASES = {
    'full-width banner / hero',
    'rating / award badges under hero',
    'testimonial carousel',
    'mid-page band',
    'below the fold',
    'below the fold, after the sitemap',
    '16pt mid-page band',
    'rotating carousel, 24pt type',
    '7pt grey, below the fold, after the sitemap',
}
_ARROW_REPLACE_RE = re.compile(
    r'[“"‘\']([^”"’\']+)[”"’\']\s*(?:→|->|=>)\s*[“"‘\']([^”"’\']+)[”"’\']',
)
_INSERT_PREFIX_RE = re.compile(
    r'^(?:'
    r'add a readable caption under the badge row:\s*'
    r'|add a caption under the badges(?: stating)?:\s*'
    r'|directly beneath the carousel, add:\s*'
    r'|append to the headline block(?: with the same treatment)?:\s*'
    r'|layout instruction:\s*'
    r'|redline:\s*'
    r')',
    re.I,
)
_INSTRUCTION_TAIL_RE = re.compile(
    r'\s+(?:Rotate the|Place that disclosure|Wording alone does not|A disclosure cannot cure).*$',
    re.I,
)
_DOUBLED_QUOTE_RE = re.compile(
    r'["“][^"”]{0,120}doubled my portfolio[^"”]{0,120}["”]'
    r'(?:\s*[—–-]\s*[^·\n|]{0,80})?'
    r'(?:\s*[·.|]\s*)?',
    re.I,
)


def _layout_notes_list(value: Any) -> list[dict[str, Any]]:
    if not isinstance(value, list):
        return []
    return [item for item in value if isinstance(item, dict)]


def format_layout_section_heading(note: dict[str, Any]) -> str:
    region = str(note.get('region') or '').strip()
    if region.startswith('slide_or_page_'):
        page = note.get('page') or region.rsplit('_', 1)[-1]
        label = f'SLIDE {page}'
    else:
        label = LAYOUT_REGION_LABELS.get(region) or region.replace('_', ' ').upper()
    font = str(note.get('font_note') or '').strip()
    excerpt = str(note.get('excerpt') or '')
    if region == 'disclosures':
        missing = 'no disclosure' in font.lower() or excerpt.lower().startswith('no disclosures')
        label = 'NO DISCLOSURES' if missing else 'DISCLOSURES'
    if region in VISUAL_PASS_FONT_NOTES and (not font or font in INGEST_FONT_NOTE_ALIASES):
        font = VISUAL_PASS_FONT_NOTES[region]
    extras = [font] if font else []
    if (
        note.get('below_fold')
        and not region.startswith('slide_or_page_')
        and not any('below the fold' in part.lower() for part in extras)
    ):
        extras.append('below the fold')
    if extras:
        return f'[{label} — {", ".join(extras)}]'
    return f'[{label}]' if label else ''


def format_layout_notes_template(notes: list[dict[str, Any]] | None) -> str:
    blocks: list[str] = []
    for note in notes or []:
        if not isinstance(note, dict):
            continue
        region = str(note.get('region') or '').strip()
        if region == 'title' and (note.get('source_url') or str(note.get('font_note') or '') in HOMEPAGE_TITLE_SKIP):
            continue
        heading = format_layout_section_heading(note)
        excerpt = normalize_space(str(note.get('excerpt') or ''))
        if heading:
            blocks.append(heading)
        if excerpt:
            blocks.append(excerpt)
    return '\n\n'.join(blocks)


def marketing_template_text(
    content: MarketingContent,
    layout_notes: list[dict[str, Any]] | None = None,
) -> str:
    notes = layout_notes if layout_notes is not None else _layout_notes_list(
        getattr(content, 'layout_notes', None)
    )
    formatted = format_layout_notes_template(notes)
    if formatted:
        return formatted
    body = str(getattr(content, 'body_text', None) or '').strip()
    if body:
        return body
    return str(getattr(content, 'title', None) or getattr(content, 'name', None) or '').strip()


def _paragraphs_from_template(text: str) -> list[str]:
    blocks = [block.strip() for block in str(text or '').replace('\r\n', '\n').split('\n\n')]
    paragraphs = [block for block in blocks if block]
    if paragraphs:
        return paragraphs
    lines = [line.strip() for line in str(text or '').split('\n') if line.strip()]
    return lines or [str(text or '').strip() or '(No extractable body text.)']


def _template_has_layout_headings(text: str) -> bool:
    return bool(re.search(r'\[[A-Z][A-Z0-9 /,-]{1,40} — ', text or ''))


def build_marketing_source_docx(
    *,
    title: str,
    format_label: str,
    body_text: str,
    rule_labels: list[str],
    extra_excerpts: list[str] | None = None,
) -> bytes:
    from docx import Document

    document = Document()
    layout_source = _template_has_layout_headings(body_text)
    if not layout_source:
        document.add_heading(title or 'Marketing content', level=1)
        if format_label:
            document.add_paragraph(f'Format: {format_label}')
        if rule_labels:
            document.add_paragraph(f'Applicable Regulation: {"; ".join(rule_labels)}')
    for block in _paragraphs_from_template(body_text):
        document.add_paragraph(normalize_space(block))
    blob = normalize_space(body_text).lower()
    for excerpt in extra_excerpts or []:
        trimmed = normalize_space(excerpt)
        if trimmed and trimmed.lower() not in blob:
            document.add_paragraph(trimmed)
    if not document.paragraphs:
        document.add_heading(title or 'Marketing content', level=1)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _layout_excerpt_for_region(region: str, notes: list[dict[str, Any]] | None) -> str:
    wanted = str(region or '').strip().lower()
    if not wanted:
        return ''
    for note in notes or []:
        name = str(note.get('region') or '').strip().lower()
        if name == wanted or name.startswith(wanted) or wanted.startswith(name):
            return normalize_space(str(note.get('excerpt') or ''))
    return ''


def _search_in_template(candidate: str, template_text: str) -> bool:
    needle = normalize_space(candidate)
    if not needle:
        return False
    blob = normalize_space(template_text)
    return needle.lower() in blob.lower()


def _is_placeholder_excerpt(excerpt: str) -> bool:
    return normalize_space(excerpt).lower() in PLACEHOLDER_EXCERPTS


def _overlap_score(left: str, right: str) -> int:
    words_left = set(re.findall(r"[a-z0-9%]{4,}", (left or '').lower()))
    words_right = set(re.findall(r"[a-z0-9%]{4,}", (right or '').lower()))
    return len(words_left & words_right)


def _resolved_region(finding: dict[str, Any], layout_notes: list[dict[str, Any]] | None) -> str:
    explicit = str(finding.get('region') or '').strip().lower()
    if explicit:
        return explicit
    finding_id = str(finding.get('id') or finding.get('finding_id') or '').strip()
    mapped = FINDING_REGION_IDS.get(finding_id)
    if mapped:
        return mapped
    blob = ' '.join(
        str(finding.get(key) or '')
        for key in ('excerpt', 'title', 'why', 'finding', 'sentence')
    ).lower()
    if 'testimonial' in blob or 'carousel' in blob:
        return 'testimonial'
    if 'badge' in blob or 'rating' in blob:
        return 'badge_row'
    if 'footer' in blob or 'below the fold' in blob:
        return 'footer'
    if 'hero' in blob or 'headline' in blob:
        return 'hero'
    best_region = ''
    best_score = 1
    excerpt = str(finding.get('excerpt') or finding.get('sentence') or '')
    for note in layout_notes or []:
        region = str(note.get('region') or '').strip().lower()
        if not region or region in SKIP_LAYOUT_REGIONS:
            continue
        score = _overlap_score(excerpt, str(note.get('excerpt') or ''))
        if score > best_score:
            best_score = score
            best_region = region
    return best_region


def _target_paragraph_for_finding(
    finding: dict[str, Any],
    *,
    template_text: str,
    layout_notes: list[dict[str, Any]] | None,
) -> str:
    excerpt = normalize_space(str(finding.get('excerpt') or ''))
    region = _resolved_region(finding, layout_notes)
    region_excerpt = _layout_excerpt_for_region(region, layout_notes)
    if region_excerpt:
        return region_excerpt
    if excerpt and not _is_placeholder_excerpt(excerpt) and _search_in_template(excerpt, template_text):
        return excerpt
    for note in layout_notes or []:
        paragraph = normalize_space(str(note.get('excerpt') or ''))
        if paragraph and excerpt and excerpt.lower() in paragraph.lower():
            return paragraph
        if paragraph and not _is_placeholder_excerpt(excerpt) and _overlap_score(excerpt, paragraph) >= 3:
            return paragraph
    sentence = normalize_space(str(finding.get('sentence') or ''))
    if excerpt and not _is_placeholder_excerpt(excerpt):
        return excerpt
    if sentence and ' — ' not in sentence:
        return sentence
    return region_excerpt


def _draft_for_finding(finding: dict[str, Any]) -> str:
    draft = str(finding.get('draft_text') or finding.get('fix') or '').strip()
    if draft.lower() not in EMPTY_DRAFTS:
        return draft
    try:
        from app.services.marketing_review_engine import _draft_text_for

        filled = str(_draft_text_for(finding, {}) or '').strip()
        if filled.lower() not in EMPTY_DRAFTS:
            return filled
    except Exception:
        filled = ''
    option = selected_or_recommended_option(finding).lower()
    if option == 'add_disclosure':
        return 'Add the required disclosure next to the claim, at a readable size.'
    if option == 'redline':
        return str(finding.get('why') or filled or '').strip()[:800]
    return draft


def _arrow_replacement(draft: str) -> tuple[str, str] | None:
    match = _ARROW_REPLACE_RE.search(draft or '')
    if not match:
        return None
    return match.group(1).strip(), match.group(2).strip()


def _publishable_insert(draft: str) -> str:
    text = str(draft or '').strip()
    if not text:
        return ''
    if _arrow_replacement(text):
        return ''
    text = _INSERT_PREFIX_RE.sub('', text).strip()
    text = _INSTRUCTION_TAIL_RE.sub('', text).strip()
    lowered = text.lower()
    if lowered.startswith('layout instruction') or lowered.startswith('recommended action'):
        return ''
    return text


def _rotate_performance_quote(text: str) -> str:
    cleaned = _DOUBLED_QUOTE_RE.sub('', text)
    cleaned = re.sub(r'\s{2,}', ' ', cleaned)
    cleaned = re.sub(r'\s+([·.|])', r' \1', cleaned)
    return cleaned.strip(' ·.|')


def _apply_arrow_to_search(search: str, old: str, new: str) -> str:
    if normalize_space(old).lower() == normalize_space(search).lower():
        return new
    index = search.lower().find(old.lower())
    if index >= 0:
        return f'{search[:index]}{new}{search[index + len(old):]}'
    return new


def _body_replacement_for_search(
    finding: dict[str, Any],
    search: str,
) -> str | None:
    option = selected_or_recommended_option(finding).lower()
    finding_id = str(finding.get('id') or finding.get('finding_id') or '').strip()
    draft = _draft_for_finding(finding)
    if option not in EDIT_OPTIONS:
        return None
    if option == 'pull':
        insert = _publishable_insert(draft)
        if insert and insert.lower() not in {
            'pull',
            'pull the claim',
            'pull the quote',
            'pull the sentence',
        }:
            return insert
        return PULL_REPLACEMENT
    arrow = _arrow_replacement(draft)
    if arrow:
        return _apply_arrow_to_search(search, arrow[0], arrow[1])
    insert = _publishable_insert(draft)
    if finding_id == 'layout_buried':
        if option != 'redline':
            return None
        insert = insert or normalize_space(draft)
        if not insert:
            return None
        if insert.lower() not in search.lower():
            return f'{search} {insert}'.strip()
        return None
    if not insert:
        return None
    region = _resolved_region(finding, None)
    if option == 'add_disclosure':
        body = search
        if 'rotate' in draft.lower() and _DOUBLED_QUOTE_RE.search(body):
            body = _rotate_performance_quote(body)
        if insert.lower() not in body.lower():
            return f'{body} {insert}'.strip()
        return body if body != search else None
    if finding_id in APPEND_REDLINE_IDS or (region == 'hero' and '%' in search):
        if insert.lower() not in search.lower():
            return f'{search} {insert}'.strip()
        return search
    if search.lower() not in insert.lower() and insert.lower() not in search.lower():
        return f'{search} {insert}'.strip()
    return insert


PRESENTATION_LAYOUT_REGIONS = frozenset({
    'slide_title',
    'bar_chart',
    'callout',
    'slide_footer',
    'appendix',
    'diagram',
    'gross_net',
    'benchmark',
    'slide_copy',
})
MAX_PRESENTATION_SEARCH = 220


def _looks_like_presentation(
    layout_notes: list[dict[str, Any]] | None,
    template_text: str = '',
) -> bool:
    for note in layout_notes or []:
        region = str(note.get('region') or '')
        if region.startswith('slide_or_page_') or region in PRESENTATION_LAYOUT_REGIONS:
            return True
    return bool(re.search(r'\[SLIDE \d+', template_text or ''))


def _clip_search_text(search: str, template_text: str) -> str:
    text = normalize_space(search)
    if not text:
        return ''
    if len(text) <= MAX_PRESENTATION_SEARCH and _search_in_template(text, template_text):
        return text
    for length in (MAX_PRESENTATION_SEARCH, 160, 120, 80):
        snippet = text[:length]
        if ' ' in snippet:
            snippet = snippet.rsplit(' ', 1)[0] or snippet
        if snippet and _search_in_template(snippet, template_text):
            return snippet
    return text[:MAX_PRESENTATION_SEARCH]


def _slide_excerpt_for_finding(
    finding: dict[str, Any],
    layout_notes: list[dict[str, Any]] | None,
) -> str:
    excerpt = normalize_space(str(finding.get('excerpt') or ''))
    page = finding.get('slide')
    try:
        page_num = int(page) if page is not None else None
    except (TypeError, ValueError):
        page_num = None
    page_candidates: list[str] = []
    for note in layout_notes or []:
        region = str(note.get('region') or '')
        paragraph = normalize_space(str(note.get('excerpt') or ''))
        if not paragraph:
            continue
        note_page = note.get('page')
        try:
            note_page_num = int(note_page) if note_page is not None else None
        except (TypeError, ValueError):
            note_page_num = None
        same_slide = False
        if page_num is not None:
            if note_page_num == page_num or region.endswith(f'_{page_num}'):
                same_slide = True
        if excerpt and not _is_placeholder_excerpt(excerpt) and excerpt.lower() in paragraph.lower():
            return paragraph
        if same_slide:
            page_candidates.append(paragraph)
    if page_candidates:
        return page_candidates[0]
    if excerpt and not _is_placeholder_excerpt(excerpt):
        return excerpt
    return ''


def _presentation_search_for_finding(
    finding: dict[str, Any],
    *,
    template_text: str,
    layout_notes: list[dict[str, Any]] | None,
) -> str:
    search = normalize_space(
        _target_paragraph_for_finding(
            finding,
            template_text=template_text,
            layout_notes=layout_notes,
        )
        or _slide_excerpt_for_finding(finding, layout_notes)
    )
    excerpt = normalize_space(str(finding.get('excerpt') or ''))
    if excerpt and not _is_placeholder_excerpt(excerpt) and excerpt.lower() in (search or '').lower():
        index = search.lower().find(excerpt.lower())
        start = max(0, index - 40)
        search = search[start: start + max(len(excerpt) + 80, MAX_PRESENTATION_SEARCH)]
    return _clip_search_text(search, template_text)


def _presentation_replace_text(finding: dict[str, Any], search: str) -> str:
    option = selected_or_recommended_option(finding).lower()
    marker = finding.get('marker') or ''
    draft = _draft_for_finding(finding)
    if option == 'pull':
        insert = _publishable_insert(draft)
        if insert and insert.lower() not in {'pull', 'pull the claim', 'pull the quote'}:
            replacement = insert if insert.lower() in search.lower() else f'{search} {insert}'.strip()
        else:
            replacement = PULL_REPLACEMENT
        if replacement == search:
            replacement = f'{search} [C{marker} pulled — do not publish this copy.]'
        return normalize_space(replacement)
    nxt = _body_replacement_for_search(finding, search)
    if nxt and nxt != search:
        return normalize_space(nxt)
    insert = _publishable_insert(draft) or normalize_space(draft)[:500]
    if insert and insert.lower() not in search.lower():
        return normalize_space(f'{search} {insert}')
    return normalize_space(f'{search} [C{marker} revision for CCO review.]')


def _presentation_replacements_from_findings(
    findings: list[dict[str, Any]],
    *,
    template_text: str,
    layout_notes: list[dict[str, Any]] | None,
) -> list[TrackedReplacement]:
    replacements: list[TrackedReplacement] = []
    used_search: set[str] = set()
    for item in findings:
        if not isinstance(item, dict):
            continue
        option = selected_or_recommended_option(item).lower()
        if option not in EDIT_OPTIONS:
            continue
        search = _presentation_search_for_finding(
            item,
            template_text=template_text,
            layout_notes=layout_notes,
        )
        if not search:
            continue
        key = search.lower()
        if key in used_search:
            search = _clip_search_text(f'{search} {item.get("id") or ""}'.strip(), template_text) or search
        used_search.add(search.lower())
        replace = _presentation_replace_text(item, search)
        if not replace or replace == search:
            replace = f'{search} [C{item.get("marker") or ""} revision for CCO review.]'
        marker = item.get('marker') or len(replacements) + 1
        region = str(item.get('anchor') or item.get('region') or 'slide copy')
        comment = format_gap_rule_fix_comment(
            rule=f'C{marker} · {str(item.get("rule") or "").strip()}'.strip(' ·'),
            gap=(
                f'Anchored to {region}. '
                + str(item.get('gap') or item.get('why') or '')
            ).strip()
            or None,
            fix=str(item.get('fix') or item.get('draft_text') or replace).strip() or None,
        )
        replacements.append(
            TrackedReplacement(
                search_text=search,
                replace_text=replace,
                comment=comment,
            )
        )
    return replacements


def tracked_replacements_from_findings(
    findings: list[dict[str, Any]],
    *,
    template_text: str = '',
    layout_notes: list[dict[str, Any]] | None = None,
) -> list[TrackedReplacement]:
    if _looks_like_presentation(layout_notes, template_text):
        return _presentation_replacements_from_findings(
            findings,
            template_text=template_text,
            layout_notes=layout_notes,
        )
    grouped: dict[str, list[dict[str, Any]]] = {}
    order: list[str] = []
    for item in findings:
        if not isinstance(item, dict):
            continue
        option = selected_or_recommended_option(item).lower()
        if option not in EDIT_OPTIONS:
            continue
        search = normalize_space(
            _target_paragraph_for_finding(
                item,
                template_text=template_text,
                layout_notes=layout_notes,
            )
        )
        if not search:
            continue
        grouped.setdefault(search, []).append(item)
        if search not in order:
            order.append(search)
    replacements: list[TrackedReplacement] = []
    for search in order:
        text = search
        comments: list[str] = []
        for item in grouped[search]:
            nxt = normalize_space(_body_replacement_for_search(item, text) or '')
            if not nxt or nxt == text:
                continue
            text = nxt
            comments.append(
                format_gap_rule_fix_comment(
                    rule=str(item.get('rule') or '').strip() or None,
                    gap=str(item.get('gap') or item.get('why') or '').strip() or None,
                    fix=str(item.get('fix') or item.get('draft_text') or text).strip() or None,
                )
            )
        if text == search:
            continue
        replacements.append(
            TrackedReplacement(
                search_text=search,
                replace_text=text,
                comment='\n\n'.join(part for part in comments if part),
            )
        )
    return replacements


def build_marketing_redline_bytes(
    *,
    title: str,
    format_label: str,
    body_text: str,
    rule_labels: list[str],
    findings: list[dict[str, Any]],
    gap_summary: str = '',
    layout_notes: list[dict[str, Any]] | None = None,
) -> bytes:
    template = format_layout_notes_template(layout_notes) or body_text
    replacements = tracked_replacements_from_findings(
        findings,
        template_text=template,
        layout_notes=layout_notes,
    )
    blob = normalize_space(template).lower()
    extras = [
        item.search_text
        for item in replacements
        if item.search_text and item.search_text.lower() not in blob
    ]
    source = build_marketing_source_docx(
        title=title,
        format_label=format_label,
        body_text=template,
        rule_labels=rule_labels,
        extra_excerpts=extras,
    )
    if not replacements:
        return source
    return apply_tracked_redline_overlay(
        source,
        replacements=replacements,
        gap_summary=gap_summary,
    )


def redline_bytes_have_rule_gap_fix(document_bytes: bytes) -> bool:
    try:
        with ZipFile(BytesIO(document_bytes)) as archive:
            names = archive.namelist()
            blob = ''
            if 'word/comments.xml' in names:
                blob = archive.read('word/comments.xml').decode('utf-8', errors='ignore')
            blob += archive.read('word/document.xml').decode('utf-8', errors='ignore')
    except Exception:
        return False
    upper = blob.upper()
    return 'RULE:' in upper and 'GAP:' in upper and 'FIX:' in upper


def _content_and_review_for_gap(
    db: Session,
    gap: GapAnalysis,
) -> tuple[MarketingContent | None, MarketingReview | None]:
    metadata = gap.retrieval_metadata if isinstance(gap.retrieval_metadata, dict) else {}
    content_id = str(metadata.get('marketing_content_id') or gap.trigger_id or '').strip()
    review_id = str(metadata.get('marketing_review_id') or '').strip()
    content = None
    review = None
    if content_id:
        content = (
            db.query(MarketingContent)
            .filter(
                MarketingContent.id == content_id,
                MarketingContent.org_id == gap.org_id,
            )
            .one_or_none()
        )
    if review_id:
        review = (
            db.query(MarketingReview)
            .filter(
                MarketingReview.id == review_id,
                MarketingReview.org_id == gap.org_id,
            )
            .one_or_none()
        )
    return content, review


def _get_or_create_pending_revision(
    db: Session,
    *,
    gap: GapAnalysis,
    policy: Policy,
    enqueue: bool = True,
) -> tuple[PolicyRevision, bool]:
    existing = (
        db.query(PolicyRevision)
        .filter(
            PolicyRevision.gap_analysis_id == gap.id,
            PolicyRevision.org_id == gap.org_id,
            PolicyRevision.status == 'PENDING',
        )
        .order_by(PolicyRevision.created_at.desc())
        .first()
    )
    if existing is not None:
        if enqueue and existing.review_requested_at is None:
            existing.review_requested_at = utc_now()
            db.flush()
            return existing, True
        return existing, False
    revision = PolicyRevision(
        id=str(uuid4()),
        gap_analysis_id=gap.id,
        policy_id=policy.id,
        org_id=gap.org_id,
        redline_file_url=None,
        status='PENDING',
        review_requested_at=utc_now() if enqueue else None,
        created_at=utc_now(),
    )
    db.add(revision)
    db.flush()
    return revision, True


def _notify_payload(revision: PolicyRevision, gap: GapAnalysis, policy: Policy, db: Session) -> dict[str, Any]:
    organization = (
        db.query(Organization)
        .filter(Organization.auth_org_id == gap.org_id)
        .one_or_none()
    )
    return {
        'revision_id': revision.id,
        'org_id': revision.org_id,
        'gap_analysis_id': gap.id,
        'policy_name': policy.name,
        'impact_score': gap.impact_score,
        'publication_id': gap.publication_id,
        'approval_chain': dict(getattr(organization, 'approval_chain', None) or {})
        if isinstance(getattr(organization, 'approval_chain', None), dict)
        else {},
        'organization_name': getattr(organization, 'name', None) or gap.org_id,
    }


def attach_marketing_redline_file(
    db: Session,
    *,
    revision: PolicyRevision,
    policy: Policy,
    redline_bytes: bytes,
) -> PolicyRevision:
    if revision.redline_file_url:
        return revision
    try:
        stored_file = upload_file(
            REDLINE_BUCKET,
            storage_path(
                policy.org_id,
                'policy-revisions',
                revision.id,
                redline_filename(policy),
            ),
            redline_bytes,
            DOCX_CONTENT_TYPE,
        )
    except Exception as exc:
        logger.warning(
            'marketing_redline_upload_failed',
            revision_id=revision.id,
            error_type=type(exc).__name__,
            error=str(exc)[:400],
        )
        return revision
    return persist_revision_redline_url(
        revision.id,
        policy.org_id,
        stored_file.uri,
        db,
    )


def build_redline_for_marketing_gap(
    db: Session,
    gap: GapAnalysis,
    content: MarketingContent,
    review: MarketingReview | None = None,
) -> bytes:
    metadata = gap.retrieval_metadata if isinstance(gap.retrieval_metadata, dict) else {}
    if review is None:
        _, review = _content_and_review_for_gap(db, gap)
    notes = _layout_notes_list(getattr(review, 'layout_notes', None) if review is not None else None)
    if not notes:
        notes = _layout_notes_list(getattr(content, 'layout_notes', None))
    if not notes:
        notes = _layout_notes_list(metadata.get('layout_notes'))
    if review is not None and isinstance(review.findings, list) and review.findings:
        from app.services.marketing_gap_bridge import apply_selected_options

        findings = apply_selected_options(
            [item for item in review.findings if isinstance(item, dict)],
            list(review.selected_options or metadata.get('selected_options') or []),
        )
    else:
        findings = [
            item
            for item in (metadata.get('sentence_findings') or [])
            if isinstance(item, dict)
        ]
    rule_labels = [str(label) for label in (metadata.get('marketing_rule_labels') or []) if str(label).strip()]
    from app.services.marketing_gap_bridge import marketing_format_label, marketing_rule_labels

    if not rule_labels:
        codes = [str(code) for code in (metadata.get('marketing_rule_codes') or [])]
        rule_labels = marketing_rule_labels(codes)
    format_label = marketing_format_label(content.format) or str(metadata.get('marketing_format') or '')
    return build_marketing_redline_bytes(
        title=str(content.title or content.name or policy_name_fallback(gap)),
        format_label=format_label,
        body_text=marketing_template_text(content, layout_notes=notes),
        rule_labels=rule_labels,
        findings=findings,
        gap_summary=str(gap.gap_summary or ''),
        layout_notes=notes,
    )


def policy_name_fallback(gap: GapAnalysis) -> str:
    return str(gap.gap_summary or 'Marketing content')[:80]


def send_marketing_review_to_cco_queue(
    db: Session,
    *,
    content: MarketingContent,
    review: MarketingReview,
    selected_options: list[dict[str, str]],
    background_tasks: Any | None = None,
    enqueue: bool = True,
    force_rebuild: bool = False,
) -> tuple[PolicyRevision | None, dict[str, Any] | None]:
    """Persist selected findings as a gap analysis and enqueue a Review Queue redline.

    Flushes only; the caller commits so a failed enqueue can be retried.
    Pass enqueue=False to prepare the same revision/redline without opening a
    CCO Review Queue item or sending reviewer notifications.
    """
    del background_tasks
    existing = (
        db.query(GapAnalysis)
        .filter(
            GapAnalysis.org_id == content.org_id,
            GapAnalysis.trigger_type == MARKETING_GAP_TRIGGER,
            GapAnalysis.trigger_id == content.id,
        )
        .order_by(GapAnalysis.created_at.desc())
        .first()
    )
    cosine = None
    if existing is not None and isinstance(existing.retrieval_metadata, dict):
        stored = existing.retrieval_metadata.get('homepage_rule_similarity')
        cosine = stored if isinstance(stored, dict) else None
    gap_id = persist_marketing_gap_analysis(
        db,
        content=content,
        review=review,
        findings=list(review.findings or []),
        cosine=cosine,
        selected_options=selected_options,
        for_cco_review=True,
    )
    if gap_id is None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail='Unable to send this marketing review to the CCO Review Queue.',
        )
    gap = (
        db.query(GapAnalysis)
        .filter(GapAnalysis.id == gap_id, GapAnalysis.org_id == content.org_id)
        .one_or_none()
    )
    if gap is None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail='Unable to send this marketing review to the CCO Review Queue.',
        )
    policy = (
        db.query(Policy)
        .filter(Policy.id == gap.policy_id, Policy.org_id == content.org_id)
        .one_or_none()
    )
    if policy is None:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail='Unable to attach a marketing redline without a content record',
        )
    revision, newly_requested = _get_or_create_pending_revision(
        db, gap=gap, policy=policy, enqueue=enqueue,
    )
    if force_rebuild and revision.redline_file_url:
        setattr(revision, 'redline_file_url', None)
        db.flush()
    if not revision.redline_file_url:
        try:
            redline_bytes = build_redline_for_marketing_gap(db, gap, content, review)
            revision = attach_marketing_redline_file(
                db,
                revision=revision,
                policy=policy,
                redline_bytes=redline_bytes,
            )
        except Exception:
            logger.exception(
                'marketing_redline_build_failed',
                revision_id=revision.id,
                gap_id=gap.id,
                content_id=content.id,
            )
    notify_payload = (
        _notify_payload(revision, gap, policy, db) if enqueue and newly_requested else None
    )
    logger.info(
        'marketing_sent_to_cco_review',
        gap_id=gap.id,
        revision_id=revision.id,
        content_id=content.id,
        redline_ready=bool(revision.redline_file_url),
        newly_requested=newly_requested,
        enqueue=enqueue,
    )
    return revision, notify_payload


async def ensure_marketing_redline_for_revision(
    revision_id: str,
    org_id: str,
    db: Session | None = None,
) -> PolicyRevisionResponse:
    from app.db.session import session_scope

    def _run(scoped_db: Session) -> PolicyRevision:
        revision = (
            scoped_db.query(PolicyRevision)
            .filter(
                PolicyRevision.id == revision_id,
                PolicyRevision.org_id == org_id,
                PolicyRevision.status == 'PENDING',
            )
            .one_or_none()
        )
        if revision is None:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail='Revision not found')
        if revision.redline_file_url:
            return revision
        gap = (
            scoped_db.query(GapAnalysis)
            .filter(
                GapAnalysis.id == revision.gap_analysis_id,
                GapAnalysis.org_id == org_id,
            )
            .one_or_none()
        )
        if gap is None or gap.trigger_type != MARKETING_GAP_TRIGGER:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail='Revision not found')
        policy = (
            scoped_db.query(Policy)
            .filter(Policy.id == revision.policy_id, Policy.org_id == org_id)
            .one_or_none()
        )
        content, review = _content_and_review_for_gap(scoped_db, gap)
        if policy is None or content is None:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail='Marketing content for this review is no longer available',
            )
        redline_bytes = build_redline_for_marketing_gap(scoped_db, gap, content, review)
        return attach_marketing_redline_file(
            scoped_db,
            revision=revision,
            policy=policy,
            redline_bytes=redline_bytes,
        )

    if db is not None:
        revision = _run(db)
        return policy_revision_response(revision)
    with session_scope(commit=False) as scoped_db:
        revision = _run(scoped_db)
        return policy_revision_response(revision)


async def generate_marketing_redline_for_gap(
    gap_analysis_id: str,
    org_id: str,
    db: Session,
    background_tasks: Any | None = None,
) -> PolicyRevisionResponse:
    gap = (
        db.query(GapAnalysis)
        .filter(GapAnalysis.id == gap_analysis_id, GapAnalysis.org_id == org_id)
        .one_or_none()
    )
    if gap is None or gap.trigger_type != MARKETING_GAP_TRIGGER:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail='Gap analysis not found')
    content, review = _content_and_review_for_gap(db, gap)
    if content is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail='Marketing content for this gap analysis is no longer available',
        )
    if review is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail='Marketing review not found')
    revision, notify_payload = send_marketing_review_to_cco_queue(
        db,
        content=content,
        review=review,
        selected_options=list(review.selected_options or []),
        background_tasks=background_tasks,
    )
    if revision is None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail='This marketing review has no findings to send to CCO review',
        )
    if notify_payload and background_tasks is not None:
        from app.core.config import get_settings
        from app.services.redline_generation_service import notify_new_revision_from_payload

        background_tasks.add_task(
            notify_new_revision_from_payload,
            notify_payload,
            get_settings(),
        )
    return policy_revision_response(revision)


def revision_is_marketing(db: Session, revision: PolicyRevision) -> bool:
    gap = (
        db.query(GapAnalysis)
        .filter(
            GapAnalysis.id == revision.gap_analysis_id,
            GapAnalysis.org_id == revision.org_id,
        )
        .one_or_none()
    )
    return gap is not None and gap.trigger_type == MARKETING_GAP_TRIGGER
