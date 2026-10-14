"""Generate regulatory analysis documents for business changes."""

from __future__ import annotations

import asyncio
import json
import re
from dataclasses import dataclass
from datetime import UTC, date, datetime
from io import BytesIO
from typing import Any

import structlog
from docx import Document
from fastapi import HTTPException, status
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from app.ai import ChatMessage, LLMProvider, is_non_retryable_error, optional_llm_provider, retry_delay
from app.core.config import Settings, get_settings
from app.core.observability import langfuse_workflow_trace, normalize_trace_tags
from app.db.session import SessionLocal, session_scope
from app.models.core import BusinessChange, Organization
from app.schemas.agent import AgentRunResponse
from app.services.llm_json import extract_json_string_field, response_model
from app.services.llm_persona import with_financial_law_expert_persona
from app.services.markdown_text import strip_markdown_emphasis
from app.services.policy_comparison_service import (
    RegulatoryReferenceAttachment,
    load_regulatory_reference_attachments,
)
from app.services.policy_revision_document import add_body_paragraphs
from app.services.storage_service import create_download_url, storage_path, upload_file

logger = structlog.get_logger(__name__)

BUSINESS_CHANGE_ANALYSIS_BUCKET = 'business-change-analyses'
DOCX_CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
ANALYSIS_STATUS_PENDING = 'PENDING'
ANALYSIS_STATUS_GENERATING = 'GENERATING'
ANALYSIS_STATUS_READY = 'READY'
ANALYSIS_STATUS_FAILED = 'FAILED'
BUSINESS_CHANGE_ANALYSIS_MAX_TOKENS = 4_096
LLM_BUSINESS_CHANGE_ANALYSIS_MAX_ATTEMPTS = 3
LLM_BUSINESS_CHANGE_ANALYSIS_RETRY_DELAY_SECONDS = 2.0
BUSINESS_CHANGE_ANALYSIS_DRAFT_TIMEOUT_SECONDS = 90.0

BUSINESS_CHANGE_ANALYSIS_SYSTEM_PROMPT = with_financial_law_expert_persona(
    """Analyze an internal business change for a US financial-services firm against
the supplied regulatory reference attachments.

Mandatory sequencing — complete before drafting recommendations:
1. Deeply read and understand the business change description (and any event summary): what
   is changing; who/what is affected; products, processes, staffing, communications, or
   books-and-records implications; and timing.
2. Deeply read the supplied REGULATORY_REFERENCE attachments (SEC, FINRA, CFTC) and extract
   the concrete obligations that apply to that change.
3. Only then synthesize executive conclusions, implications, and recommended actions.

Each attachment contains the full text of an active REGULATORY_REFERENCE policy
synced from government sources (SEC, FINRA, CFTC).
Ground every conclusion in the provided regulatory reference documents and the
business change description. Do not analyze from keywords alone.

Return only strict, valid JSON with exactly this shape (emit phase1_brief first):
{
  "phase1_brief": {
    "policy_understanding": "",
    "companion_understanding": "concise summary of the business change and applicable regs",
    "covered_well": "aspects already well bounded by existing regs/practice",
    "silent_or_thin": "areas needing careful compliance attention for this change"
  },
  "executive_summary": "concise overview of the change and regulatory relevance",
  "regulatory_considerations": [
    {
      "regulator": "FINRA",
      "topic": "short topic label",
      "analysis": "how the government reference applies to this business change",
      "source_reference": "reference policy or publication title"
    }
  ],
  "compliance_implications": ["implication 1", "implication 2"],
  "recommended_actions": ["action 1", "action 2"],
  "citations_used": [
    {
      "source": "FINRA",
      "reference": "reference policy or publication title",
      "excerpt": "short supporting excerpt"
    }
  ]
}

Keep executive_summary under 250 words. Limit regulatory_considerations to the most material items.
Do not invent regulations or citations. Escape double quotes and newlines inside JSON string values as \\n.
Do not use trailing commas or markdown fences."""
)


class RegulatoryConsideration(BaseModel):
    regulator: str = Field(default='', max_length=100)
    topic: str = Field(default='', max_length=500)
    analysis: str = Field(default='', max_length=4_000)
    source_reference: str = Field(default='', max_length=500)


class AnalysisCitation(BaseModel):
    source: str = Field(default='', max_length=100)
    reference: str = Field(default='', max_length=500)
    excerpt: str = Field(default='', max_length=4_000)


class BusinessChangeAnalysisPayload(BaseModel):
    executive_summary: str = Field(min_length=1, max_length=4_000)
    regulatory_considerations: list[RegulatoryConsideration] = Field(default_factory=list)
    compliance_implications: list[str] = Field(default_factory=list)
    recommended_actions: list[str] = Field(default_factory=list)
    citations_used: list[AnalysisCitation] = Field(default_factory=list)


@dataclass(frozen=True)
class BusinessChangeSnapshot:
    id: str
    org_id: str
    change_type: str
    description: str
    change_date: date | None
    event_summary: str
    triggers_review: bool
    assessment_rationale: str


def snapshot_from_change(change: BusinessChange) -> BusinessChangeSnapshot:
    return BusinessChangeSnapshot(
        id=change.id,
        org_id=change.org_id,
        change_type=change.change_type,
        description=change.description,
        change_date=change.change_date,
        event_summary=change.event_summary or '',
        triggers_review=bool(change.triggers_review),
        assessment_rationale=change.assessment_rationale or '',
    )


def business_change_trigger_text(change: BusinessChange | BusinessChangeSnapshot) -> str:
    parts: list[str] = []
    summary = (change.event_summary or '').strip()
    description = (change.description or '').strip()
    if summary:
        parts.append(summary)
    if description:
        parts.append(f'Original change description: {description}')
    return '\n\n'.join(parts).strip() or description


def _set_analysis_status(db: Session, change_id: str, org_id: str, status: str) -> None:
    change = (
        db.query(BusinessChange)
        .filter(BusinessChange.id == change_id, BusinessChange.org_id == org_id)
        .one_or_none()
    )
    if change is not None:
        change.analysis_status = status
        return
    logger.error(
        'business_change_analysis_status_row_missing',
        change_id=change_id,
        org_id=org_id,
        status=status,
    )


def mark_business_change_analysis_status(change_id: str, org_id: str, status: str) -> None:
    """Persist analysis status on the tenant database (background/Inngest safe)."""
    with session_scope(org_id=org_id) as db:
        _set_analysis_status(db, change_id, org_id, status)


def _load_analysis_context(
    db: Session,
    change_id: str,
    org_id: str,
) -> tuple[BusinessChangeSnapshot, str, dict[str, Any]]:
    change = (
        db.query(BusinessChange)
        .filter(BusinessChange.id == change_id, BusinessChange.org_id == org_id)
        .one_or_none()
    )
    if change is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail='Business change not found',
        )
    organization = (
        db.query(Organization)
        .filter(Organization.auth_org_id == org_id)
        .one_or_none()
    )
    organization_name = organization.name if organization is not None else org_id
    regulatory_scope = (organization.regulatory_scope if organization else {}) or {}
    return snapshot_from_change(change), organization_name, regulatory_scope


def persist_ready_business_change_analysis_document(
    change_id: str,
    org_id: str,
    *,
    analysis: dict[str, Any] | None = None,
    force: bool = False,
    mark_ready: bool = True,
) -> AgentRunResponse:
    """Store the History DOCX from existing text. Optionally mark the row Ready."""
    if SessionLocal is None:
        return AgentRunResponse(
            run_id=change_id,
            status='skipped',
            message='Business change analysis requires a configured database',
        )

    with session_scope(org_id=org_id) as scoped_db:
        existing = (
            scoped_db.query(BusinessChange)
            .filter(BusinessChange.id == change_id, BusinessChange.org_id == org_id)
            .one_or_none()
        )
        if existing is None:
            logger.error(
                'business_change_analysis_persist_not_found',
                change_id=change_id,
                org_id=org_id,
            )
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail='Business change not found',
            )
        if (
            analysis is None
            and not force
            and existing.analysis_status == ANALYSIS_STATUS_READY
            and existing.analysis_file_url
        ):
            return AgentRunResponse(
                run_id=change_id,
                status='completed',
                message='Business change regulatory analysis document already generated',
            )
        snapshot, organization_name, _regulatory_scope = _load_analysis_context(
            scoped_db,
            change_id,
            org_id,
        )
        reference_attachments = load_regulatory_reference_attachments(scoped_db, org_id)

    payload = analysis or fallback_business_change_analysis(snapshot, reference_attachments)
    docx_bytes = build_business_change_analysis_docx(
        change=snapshot,
        organization_name=organization_name,
        analysis=payload,
        reference_attachments=reference_attachments,
    )
    stored = upload_file(
        BUSINESS_CHANGE_ANALYSIS_BUCKET,
        storage_path(
            org_id,
            'business-change-analyses',
            snapshot.id,
            f'business_change_analysis_{snapshot.id}.docx',
        ),
        docx_bytes,
        DOCX_CONTENT_TYPE,
    )

    with session_scope(org_id=org_id) as scoped_db:
        change = (
            scoped_db.query(BusinessChange)
            .filter(BusinessChange.id == change_id, BusinessChange.org_id == org_id)
            .one_or_none()
        )
        if change is not None:
            change.analysis_file_url = stored.uri
            change.analysis_generated_at = datetime.now(UTC)
            if mark_ready:
                change.analysis_status = ANALYSIS_STATUS_READY
            elif change.analysis_status != ANALYSIS_STATUS_READY:
                change.analysis_status = ANALYSIS_STATUS_GENERATING
        else:
            logger.error(
                'business_change_analysis_persist_lost_row',
                change_id=change_id,
                org_id=org_id,
            )

    logger.info(
        'business_change_analysis_document_persisted',
        change_id=change_id,
        org_id=org_id,
        mark_ready=mark_ready,
        attachment_count=len(reference_attachments),
    )
    return AgentRunResponse(
        run_id=change_id,
        status='completed',
        message='Business change regulatory analysis document generated',
    )


async def generate_business_change_analysis_document(
    change_id: str,
    org_id: str,
    db: Session | None = None,
    client: LLMProvider | None = None,
) -> AgentRunResponse:
    """Analyze a business change against government references and store a DOCX report."""
    del db  # Legacy parameter retained for compatibility; sessions are always short-lived now.

    async with langfuse_workflow_trace(
        'business-change-analysis',
        user_id=org_id,
        session_id=change_id,
        tags=normalize_trace_tags('business-change-analysis'),
        metadata={'change_id': change_id, 'org_id': org_id},
        input_data={'change_id': change_id, 'org_id': org_id},
    ):
        return await _generate_business_change_analysis_document_impl(
            change_id,
            org_id,
            client=client,
        )


async def _generate_business_change_analysis_document_impl(
    change_id: str,
    org_id: str,
    *,
    client: LLMProvider | None = None,
) -> AgentRunResponse:
    if SessionLocal is None:
        return AgentRunResponse(
            run_id=change_id,
            status='skipped',
            message='Business change analysis requires a configured database',
        )

    try:
        ready = await asyncio.to_thread(
            persist_ready_business_change_analysis_document,
            change_id,
            org_id,
            mark_ready=False,
        )
        settings = get_settings()
        try:
            with session_scope(org_id=org_id) as scoped_db:
                snapshot, organization_name, regulatory_scope = _load_analysis_context(
                    scoped_db,
                    change_id,
                    org_id,
                )
                reference_attachments = load_regulatory_reference_attachments(scoped_db, org_id)
            analysis = await asyncio.wait_for(
                draft_business_change_analysis(
                    change=snapshot,
                    organization_name=organization_name,
                    regulatory_scope=regulatory_scope,
                    reference_attachments=reference_attachments,
                    settings=settings,
                    client=client,
                ),
                timeout=BUSINESS_CHANGE_ANALYSIS_DRAFT_TIMEOUT_SECONDS,
            )
            await asyncio.to_thread(
                persist_ready_business_change_analysis_document,
                change_id,
                org_id,
                analysis=analysis,
                mark_ready=False,
            )
        except Exception as exc:
            logger.warning(
                'business_change_analysis_llm_upgrade_failed',
                change_id=change_id,
                org_id=org_id,
                error_type=type(exc).__name__,
            )
        return ready
    except HTTPException:
        try:
            with session_scope(org_id=org_id) as scoped_db:
                _set_analysis_status(scoped_db, change_id, org_id, ANALYSIS_STATUS_FAILED)
        except Exception:
            logger.exception('business_change_analysis_failed_status_update', change_id=change_id)
        raise
    except Exception as exc:
        try:
            with session_scope(org_id=org_id) as scoped_db:
                _set_analysis_status(scoped_db, change_id, org_id, ANALYSIS_STATUS_FAILED)
        except Exception:
            logger.exception('business_change_analysis_failed_status_update', change_id=change_id)
        logger.exception(
            'business_change_analysis_failed',
            change_id=change_id,
            org_id=org_id,
            error_type=type(exc).__name__,
        )
        return AgentRunResponse(
            run_id=change_id,
            status='failed',
            message='Unable to generate business change regulatory analysis document',
        )


async def draft_business_change_analysis(
    *,
    change: BusinessChangeSnapshot,
    organization_name: str,
    regulatory_scope: dict[str, Any],
    reference_attachments: list[RegulatoryReferenceAttachment],
    settings: Settings,
    client: LLMProvider | None,
) -> dict[str, Any]:
    if client is None:
        client = optional_llm_provider(change.org_id, 'analysis', platform_model=settings.anthropic_opus_model)
    if client is None:
        return fallback_business_change_analysis(change, reference_attachments)

    prompt = json.dumps(
        {
            'organization': {
                'name': organization_name,
                'regulatory_scope': regulatory_scope,
            },
            'business_change': {
                'id': change.id,
                'change_type': change.change_type,
                'description': change.description,
                'effective_date': change.change_date.isoformat() if change.change_date else None,
                'event_summary': change.event_summary,
                'triggers_review': change.triggers_review,
                'assessment_rationale': change.assessment_rationale,
            },
            'regulatory_reference_attachments': [
                attachment.as_prompt_dict() for attachment in reference_attachments
            ],
        },
        ensure_ascii=True,
    )

    last_error: Exception | None = None
    for attempt in range(1, LLM_BUSINESS_CHANGE_ANALYSIS_MAX_ATTEMPTS + 1):
        try:
            response = await client.generate(
                [ChatMessage('user', prompt)],
                max_tokens=BUSINESS_CHANGE_ANALYSIS_MAX_TOKENS,
                system=BUSINESS_CHANGE_ANALYSIS_SYSTEM_PROMPT,
                cache_system=True,
                throttle='opus',
                trace_name='business_change.analysis',
                trace_metadata={'business_change_id': change.id},
            )
            payload = response_model(
                response,
                BusinessChangeAnalysisPayload,
                loose_parser=parse_business_change_analysis_loosely,
                recovery_log_event='business_change_analysis_json_recovered_with_loose_parser',
            )
            return payload.model_dump()
        except Exception as exc:
            last_error = exc
            logger.warning(
                'business_change_analysis_drafting_failed',
                change_id=change.id,
                attempt=attempt,
                error_type=type(exc).__name__,
            )
            if is_non_retryable_error(exc):
                break
            if attempt < LLM_BUSINESS_CHANGE_ANALYSIS_MAX_ATTEMPTS:
                delay = retry_delay(exc, attempt, settings)
                if delay > 0:
                    await asyncio.sleep(delay)

    logger.warning(
        'business_change_analysis_drafting_exhausted_retries',
        change_id=change.id,
        error_type=type(last_error).__name__ if last_error else None,
    )
    return fallback_business_change_analysis(change, reference_attachments)


def fallback_business_change_analysis(
    change: BusinessChangeSnapshot,
    reference_attachments: list[RegulatoryReferenceAttachment],
) -> dict[str, Any]:
    considerations = [
        {
            'regulator': attachment.regulator_source,
            'topic': attachment.policy_name,
            'analysis': attachment.content[:1_000],
            'source_reference': attachment.policy_name,
        }
        for attachment in reference_attachments[:6]
    ]
    return {
        'executive_summary': (
            (change.event_summary or change.description).strip()
            or 'Business change logged for CCO review.'
        ),
        'regulatory_considerations': considerations,
        'compliance_implications': [
            'Review attached regulatory reference policies against the described business change.',
            'Determine whether internal policies require amendment before the change goes live.',
        ],
        'recommended_actions': [
            'Review the attached regulatory reference documents with compliance and legal.',
            'Update affected internal policies if supervisory requirements apply.',
        ],
        'citations_used': [
            {
                'source': attachment.regulator_source,
                'reference': attachment.policy_name,
                'excerpt': attachment.content[:500],
            }
            for attachment in reference_attachments[:6]
        ],
    }


def parse_business_change_analysis_loosely(text: str) -> dict[str, Any] | None:
    executive_summary = extract_json_string_field(text, ('executive_summary',))
    if not executive_summary:
        return None
    payload: dict[str, Any] = {
        'executive_summary': executive_summary,
        'regulatory_considerations': [],
        'compliance_implications': [],
        'recommended_actions': [],
        'citations_used': [],
    }

    for field in ('compliance_implications', 'recommended_actions'):
        match = re.search(rf'"{field}"\s*:\s*\[(.*?)\]', text, re.DOTALL)
        if match is not None:
            payload[field] = re.findall(r'"([^"\\]*(?:\\.[^"\\]*)*)"', match.group(1))

    return payload


def build_business_change_analysis_docx(
    *,
    change: BusinessChangeSnapshot,
    organization_name: str,
    analysis: dict[str, Any],
    reference_attachments: list[RegulatoryReferenceAttachment],
) -> bytes:
    document = Document()
    document.add_heading('Business Change Regulatory Analysis', level=0)
    document.add_paragraph(f'Organization: {organization_name}')
    document.add_paragraph(f'Generated: {datetime.now(UTC).strftime("%Y-%m-%d %H:%M UTC")}')
    document.add_paragraph(
        'This document analyzes an internal business change against active regulatory reference '
        'policies synced from government sources. It is decision support for a CCO and not legal advice.'
    )

    document.add_heading('Business Change', level=1)
    document.add_paragraph(f'Type: {change.change_type.replace("_", " ").title()}')
    if change.change_date is not None:
        document.add_paragraph(f'Effective date: {change.change_date.isoformat()}')
    document.add_paragraph(change.description.strip())

    document.add_heading('Executive Summary', level=1)
    add_body_paragraphs(document, strip_markdown_emphasis(str(analysis.get('executive_summary') or '')))

    considerations = analysis.get('regulatory_considerations') or []
    if considerations:
        document.add_heading('Regulatory Considerations', level=1)
        for item in considerations:
            if not isinstance(item, dict):
                continue
            heading = ' — '.join(
                part for part in (
                    strip_markdown_emphasis(str(item.get('regulator') or '')),
                    strip_markdown_emphasis(str(item.get('topic') or '')),
                ) if part
            )
            if heading:
                document.add_heading(heading, level=2)
            analysis_text = strip_markdown_emphasis(str(item.get('analysis') or ''))
            if analysis_text:
                add_body_paragraphs(document, analysis_text)
            source_reference = strip_markdown_emphasis(str(item.get('source_reference') or ''))
            if source_reference:
                document.add_paragraph(f'Source: {source_reference}')

    implications = analysis.get('compliance_implications') or []
    if implications:
        document.add_heading('Compliance Implications', level=1)
        for item in implications:
            cleaned = strip_markdown_emphasis(str(item))
            if cleaned:
                document.add_paragraph(cleaned, style='List Bullet')

    actions = analysis.get('recommended_actions') or []
    if actions:
        document.add_heading('Recommended Actions', level=1)
        for index, item in enumerate(actions, start=1):
            cleaned = strip_markdown_emphasis(str(item))
            if cleaned:
                document.add_paragraph(f'{index}. {cleaned}')

    document.add_heading('Regulatory Reference Attachments Reviewed', level=1)
    if reference_attachments:
        for attachment in reference_attachments:
            document.add_heading(
                f'{attachment.regulator_source}: {attachment.policy_name}',
                level=2,
            )
            document.add_paragraph(attachment.content.strip())
    else:
        document.add_paragraph(
            'No active REGULATORY_REFERENCE policies were available for this organization.'
        )

    citations = analysis.get('citations_used') or []
    if citations:
        document.add_heading('Citations Used In Analysis', level=1)
        for citation in citations:
            if not isinstance(citation, dict):
                continue
            label = strip_markdown_emphasis(
                str(citation.get('reference') or citation.get('source') or 'Reference')
            )
            document.add_heading(str(label), level=2)
            excerpt = strip_markdown_emphasis(str(citation.get('excerpt') or ''))
            if excerpt:
                add_body_paragraphs(document, excerpt)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def business_change_analysis_download(
    change: BusinessChange,
    *,
    expires_in: int = 900,
) -> dict[str, Any]:
    if change.analysis_status != ANALYSIS_STATUS_READY or not change.analysis_file_url:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail='Business change analysis document is not ready',
        )
    return {
        'change_id': change.id,
        'analysis_status': change.analysis_status,
        'file_url': change.analysis_file_url,
        'download_url': create_download_url(change.analysis_file_url, expires_in),
        'expires_in': expires_in,
        'generated_at': change.analysis_generated_at,
    }


def format_change_type(change_type: str) -> str:
    return change_type.replace('_', ' ').title()
