"""Build downloadable policy revision DOCX files using the enterprise policy template."""

from __future__ import annotations

import json
import re
from copy import deepcopy
from datetime import date
from typing import Any

from docx import Document
from docx.shared import Inches, Pt

from app.core.time import utc_now
from app.services.llm_persona import with_internal_policy_analysis_persona
from app.services.redline_document import PolicyParagraph

DEFAULT_POLICY_OWNER = "Chief Compliance Officer / Chief Operating Officer"
DEFAULT_POLICY_APPROVER = "Sarah Mitchell, Chief Compliance Officer"

CANONICAL_POLICY_SECTION_SPECS: list[dict[str, Any]] = [
    {"number": "1", "title": "Purpose", "subsections": []},
    {"number": "2", "title": "Scope", "subsections": []},
    {
        "number": "3",
        "title": "Vendor Onboarding and Due Diligence",
        "subsections": [{"number": "3.2", "title": "Contract Requirements"}],
    },
    {
        "number": "4",
        "title": "Ongoing Vendor Monitoring",
        "subsections": [
            {"number": "4.1", "title": "Periodic Reviews"},
            {"number": "4.2", "title": "Incident and Change Notification"},
        ],
    },
    {"number": "5", "title": "Vendor Termination", "subsections": []},
    {"number": "6", "title": "Roles and Responsibilities", "subsections": []},
    {"number": "7", "title": "Recordkeeping", "subsections": []},
]

ENTERPRISE_POLICY_TEMPLATE_TEXT = """<Company Name>
<Company Department>
POLICY: <Policy Name>

Policy Number: <Policy Number>
Version: <Version Number>
Effective Date: <today's date>
Owner: <Chief Compliance Officer / Chief Operating Officer>
Applicable Regulation: <Applicable Regulation>
Regulatory Oversight Report
Next Review Date: <two years from today's date>

1. Purpose

<text>

2. Scope
<text>

3. Vendor Onboarding and Due Diligence
<text>

3.2 Contract Requirements
<text>

4. Ongoing Vendor Monitoring

4.1 Periodic Reviews
<text>


4.2 Incident and Change Notification
<text>

5. Vendor Termination
<text>

6. Roles and Responsibilities
<text>

7. Recordkeeping
<text>

Approved by: Sarah Mitchell, Chief Compliance Officer
Approved date: <today's date>"""


def format_policy_display_date(value: date | None = None) -> str:
    when = value or utc_now().date()
    return f"{when.strftime('%B')} {when.day}, {when.year}"


def add_policy_years(value: date, years: int) -> date:
    try:
        return value.replace(year=value.year + years)
    except ValueError:
        return value.replace(year=value.year + years, day=28)


def policy_template_dates(reference: date | None = None) -> dict[str, str]:
    today = reference or utc_now().date()
    formatted_today = format_policy_display_date(today)
    next_review = format_policy_display_date(add_policy_years(today, 2))
    return {
        "effective_date": formatted_today,
        "next_review_date": next_review,
        "approved_date": formatted_today,
    }


def bump_policy_version_label(current: str, *, fallback: str | None = None) -> str:
    """Increment the minor component of a policy version label (e.g. 1.2 -> 1.3, v1.2 -> v1.3)."""
    source = re.sub(r"\s+", " ", (current or fallback or "").strip())
    if not source:
        return "1.1"

    source = re.sub(r"^(?:version|ver\.?)\s+", "", source, flags=re.IGNORECASE).strip()
    prefix = ""
    version_prefix = re.match(r"^([vV])(.+)$", source)
    if version_prefix:
        prefix = version_prefix.group(1)
        source = version_prefix.group(2).strip()

    match = re.match(r"^(\d+)\.(\d+)(?:\.\d+)?$", source)
    if match:
        return f"{prefix}{match.group(1)}.{int(match.group(2)) + 1}"
    if re.fullmatch(r"\d+", source):
        return f"{prefix}{int(source) + 1}"

    embedded = re.search(r"(\d+)\.(\d+)(?:\.\d+)?", source)
    if embedded:
        return f"{prefix}{embedded.group(1)}.{int(embedded.group(2)) + 1}"

    return (current or fallback or "1.1").strip() or "1.1"


def canonical_policy_section_outline() -> list[dict[str, Any]]:
    return json.loads(json.dumps(CANONICAL_POLICY_SECTION_SPECS))


def _normalize_section_title(title: str) -> str:
    return re.sub(r"\s+", " ", title.strip().lower())


def _section_has_content(section: dict[str, Any]) -> bool:
    body = str(section.get("body") or "").strip()
    if body:
        return True
    for subsection in section.get("subsections") or []:
        if isinstance(subsection, dict) and str(subsection.get("body") or "").strip():
            return True
    return False


def align_policy_sections_to_template(sections: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Map drafted section content onto the fixed enterprise policy section scaffold."""
    by_number = {
        str(section.get("number") or "").strip(): section
        for section in sections
        if str(section.get("number") or "").strip()
    }
    by_title = {
        _normalize_section_title(str(section.get("title") or "")): section
        for section in sections
        if str(section.get("title") or "").strip()
    }

    aligned: list[dict[str, Any]] = []
    for spec in CANONICAL_POLICY_SECTION_SPECS:
        source = by_number.get(spec["number"]) or by_title.get(_normalize_section_title(spec["title"]))
        subsections: list[dict[str, Any]] = []
        source_subsections = source.get("subsections") if isinstance(source, dict) else []
        if not isinstance(source_subsections, list):
            source_subsections = []

        for sub_spec in spec.get("subsections") or []:
            matched = None
            for candidate in source_subsections:
                if not isinstance(candidate, dict):
                    continue
                candidate_number = str(candidate.get("number") or "").strip()
                candidate_title = _normalize_section_title(str(candidate.get("title") or ""))
                if (
                    candidate_number == sub_spec["number"]
                    or candidate_title == _normalize_section_title(sub_spec["title"])
                ):
                    matched = candidate
                    break
            subsections.append(
                {
                    "number": sub_spec["number"],
                    "title": sub_spec["title"],
                    "body": str((matched or {}).get("body") or "").strip(),
                }
            )

        aligned.append(
            {
                "number": spec["number"],
                "title": spec["title"],
                "body": str((source or {}).get("body") or "").strip(),
                "subsections": subsections,
            }
        )

    if not any(_section_has_content(section) for section in aligned):
        for index, spec in enumerate(CANONICAL_POLICY_SECTION_SPECS):
            if index >= len(sections):
                break
            source = sections[index]
            if not isinstance(source, dict):
                continue
            aligned[index]["body"] = str(source.get("body") or source.get("revised_text") or "").strip()
            source_subsections = source.get("subsections")
            if isinstance(source_subsections, list):
                for sub_index, subsection in enumerate(source_subsections):
                    if not isinstance(subsection, dict):
                        continue
                    target_subsections = aligned[index]["subsections"]
                    if sub_index < len(target_subsections):
                        target_subsections[sub_index]["body"] = str(subsection.get("body") or "").strip()

    return aligned


def build_redline_generation_system_prompt() -> str:
    citation_example = (
        "RULE: The recordkeeping and supervision obligations above follow the "
        "communication, not the app.\\n\\n"
        "GAP: Once texting is approved, client-initiated contact on unapproved apps "
        "becomes the main leakage path for un-captured business records.\\n\\n"
        "FIX: Redirect-and-report duty keeps the record trail intact and evidences "
        "supervision of off-channel contact."
    )
    return with_internal_policy_analysis_persona(
        f"""You are drafting tracked Internal Policy redline edits (Reg AI Agent 3).

Compare the INTERNAL enterprise policy against matched REGULATORY_REFERENCE material for the SAME
regulator type only (FINRA→FINRA Regulations, SEC→SEC Regulations, CFTC→CFTC Regulations).

Phase 2 — Sentence-by-sentence redline drafting (required only after Phase 1A policy understanding
and Phase 1B companion-source understanding). Hold the holistic meaning of the full Internal Policy
and any business_change / marketing_content / regulatory / gap_analysis companions while drafting
each sentence edit:
- Prefer edits that fit the existing control design from Phase 1A.
- Do not treat a sentence as an isolated keyword string when another section already completely
  satisfies the same obligation.
- Still produce one edit per deficient sentence — do not collapse multiple sentence issues into one
  theme-level rewrite.

Critical constraints:
1. Analyze EVERY sentence in internal_policy.paragraphs one at a time. Those paragraphs start at the
   first numbered section or subtitle (for example "1. Purpose") — cover-page metadata is omitted
   on purpose and must not be invented or edited.
2. NEVER revise cover-page or policy metadata fields yourself. After drafting, the system updates:
   - Version: bump one minor step (e.g. v1.2 → v1.3)
   - Effective Date: today's date
   - Applicable Regulation: only compact rule name + number labels from each RULE comment
     (e.g. FINRA Rule 3310; SEC Rule 17a-4) — never full RULE prose
   - Next Review Date: two years from today's date
   - Approved date / Approval date / Date approved (or similar): today's date
   Put clear, concrete regulation citations in each RULE line (e.g. "FINRA Rule 3310") so Applicable
   Regulation can collect them. Do not invent Version/Effective Date/Next Review Date/Approved date
   values.
3. Preserve the original document structure and formatting. Do NOT change section titles, subsection
   titles, numbering, headings, or layout. Only revise sentence/body wording that has a discrepancy.
4. For each sentence that contradicts, omits, or otherwise conflicts with the matching regulation (or
   with a supplied business_change / marketing obligation reflected in gap findings), produce one
   edit with a Word comment containing RULE, GAP, and FIX for that sentence.
5. One paragraph may receive several edits when multiple sentences need different fixes.
6. GAP and FIX comment text must be tailored to each sentence. The same RULE citation MAY appear on
   multiple edits when distinct sentences implicate that rule. Never collapse multiple sentence edits
   into one.
7. Prefer gap_analysis_context.sentence_findings when present: each finding becomes one edit anchored
   on finding.sentence with a FIX that remediates that gap. citation_comment must use that finding's
   full RULE/GAP/FIX text (do not truncate with ellipses). Skip any finding whose sentence is
   cover-page metadata.
8. Do not invent regulations. Ground edits only in matched_regulatory_reference_material and supplied
   findings.

For each edit:
- anchor_text: verbatim sentence (or minimal contiguous phrase) from the internal policy
- replacement_text: revised policy sentence that CLOSES the gap (never empty; never delete without replace)
- citation_comment: explain WHY the change is made, as a Word comment with this exact shape:
  RULE: <complete governing obligation — not truncated>
  <blank line>
  GAP: <complete explanation of what the current language fails to do — not truncated>
  <blank line>
  FIX: <complete explanation of what the replacement achieves — not truncated>
  GAP and FIX wording must be unique per edit; the same RULE citation may be reused.
  Never end GAP or FIX with "…" / "..." — include the full sentences.

Return only strict, valid JSON with exactly this shape (emit phase1_brief before edits):
{{
  "phase1_brief": {{
    "policy_understanding": "concise summary of the Internal Policy purpose, scope, and control design",
    "companion_understanding": (
      "concise summary of regulatory references, gap findings, "
      "and any business_change or marketing context used"
    ),
    "covered_well": "what the policy already addresses adequately",
    "silent_or_thin": "where sentence-level edits are needed based on Phase 1A/1B"
  }},
  "drafting_summary": "brief summary of how many sentence-level gaps were addressed",
  "edits": [
    {{
      "anchor_text": "exact internal policy sentence to revise",
      "replacement_text": "revised sentence that remediates the discrepancy",
      "citation_comment": "{citation_example}"
    }}
  ],
  "regulatory_sources_used": [
    {{
      "policy_name": "FINRA Regulations",
      "regulator": "FINRA",
      "excerpt": "short excerpt supporting the revisions"
    }}
  ]
}}

If there are no sentence discrepancies, return "edits": [] and a drafting_summary stating no changes.
Escape double quotes and newlines inside JSON string values as \\n.
Do not use trailing commas or markdown fences.
Do not return policy_sections, policy_header, or a full rewritten template document.
Do not delete policy lines without providing replacement_text for each edit.
"""
    )


LIST_BULLET_STYLE = "List Bullet"
LIST_NUMBER_STYLE = "List Number"
LIST_MARKER_PREFIX = re.compile(r"^[\-*\u2022\u2023\u25E6\u00B7\u2219]\s+")
NUMBERED_LINE_PREFIX = re.compile(r"^\d+\.\s+")


def normalize_body_text_for_lists(body: str) -> str:
    """Split dense inline markdown bullets onto separate lines for DOCX rendering."""
    normalized = (body or "").replace("\r\n", "\n").strip()
    if not normalized:
        return ""

    normalized = re.sub(r"(?<!\*)\s\*(?!\*)\s+", "\n* ", normalized)
    normalized = re.sub(r"(?<=\S)\s+\-\s+(?=\S)", "\n- ", normalized)
    normalized = re.sub(r"\s•\s+", "\n• ", normalized)
    normalized = re.sub(r"(?<=\S)\s(\d+\.\s+)", r"\n\1", normalized)
    return normalized


def strip_list_marker(text: str) -> str:
    """Remove bullet or numbering prefixes so Word list styles provide the marker."""
    cleaned = text.strip()
    previous = None
    while cleaned and cleaned != previous:
        previous = cleaned
        cleaned = LIST_MARKER_PREFIX.sub("", cleaned, count=1).strip()
        cleaned = NUMBERED_LINE_PREFIX.sub("", cleaned, count=1).strip()
    return cleaned


def _apply_list_paragraph_formatting(paragraph: Any, *, style_name: str) -> None:
    """Copy numbering metadata and apply a hanging indent for reliable bullet rendering."""
    style = paragraph.part.document.styles[style_name]
    style_pPr = style.element.pPr
    if style_pPr is not None and style_pPr.numPr is not None:
        pPr = paragraph._p.get_or_add_pPr()
        if pPr.numPr is None:
            pPr.insert(0, deepcopy(style_pPr.numPr))

    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.first_line_indent = Inches(-0.25)
    paragraph.paragraph_format.space_after = Pt(4)


def _add_bullet_paragraph(document: Document, text: str) -> None:
    content = strip_list_marker(text)
    if not content:
        return
    paragraph = document.add_paragraph(content, style=LIST_BULLET_STYLE)
    _apply_list_paragraph_formatting(paragraph, style_name=LIST_BULLET_STYLE)


def _add_numbered_paragraph(document: Document, text: str) -> None:
    content = strip_list_marker(text)
    if not content:
        return
    paragraph = document.add_paragraph(content, style=LIST_NUMBER_STYLE)
    _apply_list_paragraph_formatting(paragraph, style_name=LIST_NUMBER_STYLE)


def _is_bullet_line(text: str) -> bool:
    return bool(LIST_MARKER_PREFIX.match(text))


def _is_numbered_line(text: str) -> bool:
    return bool(NUMBERED_LINE_PREFIX.match(text))


def add_body_paragraphs(document: Document, body: str) -> None:
    """Render policy body text with Word list styles for bullets and numbered lines."""
    normalized = normalize_body_text_for_lists(body)
    if not normalized:
        return

    for line in normalized.splitlines():
        text = line.strip()
        if not text:
            continue
        if _is_bullet_line(text):
            _add_bullet_paragraph(document, text)
            continue
        if _is_numbered_line(text):
            _add_numbered_paragraph(document, text)
            continue
        document.add_paragraph(text, style="Normal")


def sections_from_policy_paragraphs(
    paragraphs: list[PolicyParagraph],
    *,
    gap_summary: str,
) -> list[dict[str, str]]:
    """Fallback legacy sections when LLM drafting is unavailable."""
    sections: list[dict[str, str]] = []
    for item in paragraphs:
        revised = item.text.strip()
        if not revised:
            continue
        if not revised.endswith("."):
            revised = (
                f"{revised} This section must reflect applicable supervisory requirements "
                "identified in the current regulatory review."
            )
        sections.append(
            {
                "heading": item.heading,
                "original_text": item.text,
                "revised_text": revised[:4_000],
                "change_note": "Automated fallback revision pending CCO review.",
            }
        )
    if not sections:
        sections.append(
            {
                "heading": "Policy Update",
                "original_text": "",
                "revised_text": gap_summary[:2_000],
                "change_note": "Fallback draft generated from gap analysis summary.",
            }
        )
    return sections


def default_policy_header_from_context(
    *,
    company_name: str,
    policy_name: str,
    policy_version: int,
    applicable_regulation: str,
    reference_date: date | None = None,
) -> dict[str, str]:
    dates = policy_template_dates(reference_date)
    return {
        "company_name": company_name,
        "company_department": "Compliance & Legal Department",
        "policy_name": policy_name,
        "policy_number": "TBD",
        "version": str(policy_version),
        "effective_date": dates["effective_date"],
        "owner": DEFAULT_POLICY_OWNER,
        "applicable_regulation": applicable_regulation or "TBD",
        "next_review_date": dates["next_review_date"],
    }


def template_sections_from_paragraphs(paragraphs: list[PolicyParagraph]) -> list[dict[str, Any]]:
    """Build numbered template sections from extracted DOCX paragraphs."""
    sections: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    section_index = 0

    for item in paragraphs:
        text = item.text.strip()
        if not text:
            continue
        heading = (item.heading or "").strip()
        if heading:
            section_index += 1
            current = {
                "number": str(section_index),
                "title": heading,
                "body": text,
                "subsections": [],
            }
            sections.append(current)
            continue
        if current is None:
            section_index += 1
            current = {
                "number": str(section_index),
                "title": f"Section {section_index}",
                "body": text,
                "subsections": [],
            }
            sections.append(current)
        else:
            current["body"] = f"{current['body']}\n{text}".strip()

    return sections
