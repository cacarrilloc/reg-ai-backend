from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.ai import AIProviderError
from app.core.config import Settings
from app.main import app
from app.services import document_ingestion
from app.services.document_ingestion import (
    DocumentSection,
    ParsedPolicyDocument,
    TextChunk,
    canonical_document_content_type,
    chunk_policy_document,
)

DOCX_MIME_TYPE = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'


def policy_docx_bytes() -> bytes:
    document = Document()
    document.add_heading('Written Supervisory Procedures', level=1)
    document.add_paragraph(
        'The firm reviews new products, customer communications, and supervisory controls before launch.'
    )
    document.add_paragraph(
        'Compliance maintains evidence of approvals, policy exceptions, and required remediation activities.'
    )
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_upload_policy_docx_parses_and_chunks_without_openai_key() -> None:
    client = TestClient(app)
    response = client.post(
        '/api/v1/policies/upload',
        files={'file': ('finra-internal.docx', policy_docx_bytes(), DOCX_MIME_TYPE)},
        data={'policy_name': 'FINRA'},
        headers={'x-org-id': 'test-org'},
    )

    assert response.status_code == 202
    payload = response.json()
    assert payload['org_id'] == 'test-org'
    assert payload['filename'] == 'finra-internal.docx'
    assert payload['file_type'] == 'docx'
    assert payload['chunk_count'] >= 1
    assert payload['embedded_chunk_count'] == 0
    assert payload['status'] == 'parsed_without_embeddings'
    assert payload['persisted'] is False


def test_upload_policy_requires_regulator_policy_name() -> None:
    client = TestClient(app)
    response = client.post(
        '/api/v1/policies/upload',
        files={'file': ('written-supervisory-procedures.docx', policy_docx_bytes(), DOCX_MIME_TYPE)},
        headers={'x-org-id': 'test-org'},
    )

    assert response.status_code == 422


def test_upload_policy_accepts_free_text_policy_name() -> None:
    client = TestClient(app)
    response = client.post(
        '/api/v1/policies/upload',
        files={'file': ('written-supervisory-procedures.docx', policy_docx_bytes(), DOCX_MIME_TYPE)},
        data={'policy_name': 'AML Policy'},
        headers={'x-org-id': 'test-org'},
    )

    assert response.status_code == 202
    assert response.json()['filename'] == 'written-supervisory-procedures.docx'


def test_policy_name_check_normalizes_name_without_database() -> None:
    client = TestClient(app)
    response = client.get(
        '/api/v1/policies/name-check',
        params={'name': '  Written   Supervisory Procedures  '},
        headers={'x-org-id': 'test-org'},
    )

    assert response.status_code == 200
    assert response.json() == {
        'normalized_name': 'Written Supervisory Procedures',
        'exact_match': None,
        'similar_names': [],
    }


def test_upload_policy_rejects_legacy_doc_files() -> None:
    client = TestClient(app)
    response = client.post(
        '/api/v1/policies/upload',
        files={'file': ('legacy-policy.doc', b'legacy word bytes', 'application/msword')},
        data={'policy_name': 'SEC'},
        headers={'x-org-id': 'test-org'},
    )

    assert response.status_code == 415
    assert 'Legacy .doc files are not supported' in response.json()['detail']


def test_policy_document_chunking_splits_large_text_with_overlap() -> None:
    parsed = ParsedPolicyDocument(
        filename='policy.pdf',
        document_type='pdf',
        text='A' * 900,
        sections=[DocumentSection(label='page 1', index=1, text='A' * 900)],
    )

    chunks = chunk_policy_document(parsed, chunk_size=300, chunk_overlap=50)

    assert len(chunks) > 1
    assert chunks[0].metadata['filename'] == 'policy.pdf'
    assert chunks[0].metadata['section_labels'] == ['page 1']
    assert chunks[0].content[-50:] in chunks[1].content


def test_document_content_type_is_normalized_after_validation() -> None:
    assert canonical_document_content_type('pdf') == 'application/pdf'
    assert canonical_document_content_type('docx') == DOCX_MIME_TYPE


@pytest.mark.asyncio
async def test_local_upload_falls_back_when_openai_embeddings_fail(monkeypatch) -> None:
    class FailingEmbedder:
        model_label = 'text-embedding-3-large'

        async def embed(self, *_args, **_kwargs):
            raise AIProviderError('quota unavailable', provider='openai')

    monkeypatch.setattr(document_ingestion, 'optional_embedding_provider', lambda _org_id: FailingEmbedder())
    settings = Settings(ENVIRONMENT='local', OPENAI_API_KEY='test-key')
    chunks = [TextChunk(index=0, content='policy text', metadata={})]

    embeddings = await document_ingestion.generate_embeddings(chunks, settings)

    assert embeddings == []
