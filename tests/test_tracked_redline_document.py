from io import BytesIO

from docx import Document

from app.services.tracked_redline_document import (
    ParagraphOverlay,
    TrackedReplacement,
    accept_all_tracked_changes,
    apply_paragraph_overlays,
    apply_tracked_redline_overlay,
    docx_has_track_changes,
    enable_track_revisions,
    replacements_from_revision_diff,
)


def sample_internal_docx() -> bytes:
    document = Document()
    document.add_heading("1. FINRA Policy", level=1)
    document.add_paragraph(
        "Supervisors review retained electronic communications every month and document exceptions."
    )
    document.add_paragraph(
        "Employees may not use personal devices for firm business unless approved in writing."
    )
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def sample_revision_docx() -> bytes:
    document = Document()
    document.add_heading("1. FINRA Policy", level=1)
    document.add_paragraph(
        "Supervisors review retained electronic communications every month "
        "and document exceptions and retain approval evidence."
    )
    document.add_paragraph(
        "Employees may not use personal devices for firm business unless approved in writing."
    )
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_enable_track_revisions_sets_settings_flag() -> None:
    document = Document()
    enable_track_revisions(document)
    xml = document.settings.element.xml
    assert "w:trackRevisions" in xml


def test_replacements_from_revision_diff_finds_changed_paragraph() -> None:
    internal = sample_internal_docx()
    revision = sample_revision_docx()
    replacements = replacements_from_revision_diff(internal, revision)
    assert any(
        "document exceptions." in replacement.search_text
        and "approval evidence" in replacement.replace_text
        for replacement in replacements
    )


def test_apply_tracked_redline_overlay_writes_ins_and_del() -> None:
    internal = sample_internal_docx()
    output = apply_tracked_redline_overlay(
        internal,
        replacements=[
            TrackedReplacement(
                search_text=(
                    "Supervisors review retained electronic communications every month "
                    "and document exceptions."
                ),
                replace_text=(
                    "Supervisors review retained electronic communications every month "
                    "and document exceptions and retain approval evidence."
                ),
            )
        ],
        gap_summary="FINRA requires approval evidence.",
    )
    assert docx_has_track_changes(output) is True

    clean = accept_all_tracked_changes(output)
    assert docx_has_track_changes(clean) is False
    document = Document(BytesIO(clean))
    assert any("approval evidence" in paragraph.text for paragraph in document.paragraphs)


def test_apply_paragraph_overlays_replaces_whole_paragraph_with_track_changes() -> None:
    internal = sample_internal_docx()
    original = (
        "Supervisors review retained electronic communications every month "
        "and document exceptions."
    )
    revised = (
        "Supervisors review retained electronic communications every month "
        "and document exceptions and retain approval evidence."
    )
    output = apply_paragraph_overlays(
        internal,
        [
            ParagraphOverlay(
                original_text=original,
                revised_text=revised,
            )
        ],
        gap_summary="FINRA requires approval evidence.",
    )
    assert docx_has_track_changes(output) is True
    result = Document(BytesIO(output))
    assert "FINRA Policy" in result.paragraphs[0].text


def test_apply_paragraph_overlays_attaches_unique_rule_gap_fix_comments() -> None:
    document = Document()
    document.add_heading("1. Supervision", level=1)
    document.add_paragraph("First original control sentence remains incomplete.")
    document.add_paragraph("Second original control sentence lacks reporting.")
    buffer = BytesIO()
    document.save(buffer)
    internal = buffer.getvalue()
    first_comment = (
        "RULE: FINRA Rule 3110 requires written supervisory procedures.\n\n"
        "GAP: The policy omits monthly exception documentation.\n\n"
        "FIX: Require supervisors to retain approval evidence each month."
    )
    second_comment = (
        "RULE: FINRA Rule 3110(b) requires escalation of red flags.\n\n"
        "GAP: The policy does not require reporting of supervision failures.\n\n"
        "FIX: Require escalation reporting for exceptions within one business day."
    )
    output = apply_paragraph_overlays(
        internal,
        [
            ParagraphOverlay(
                original_text="First original control sentence remains incomplete.",
                revised_text="First revised control sentence documents monthly exceptions.",
                comment=first_comment,
            ),
            ParagraphOverlay(
                original_text="Second original control sentence lacks reporting.",
                revised_text="Second revised control sentence requires escalation reporting.",
                comment=second_comment,
            ),
        ],
        gap_summary="",
    )
    assert docx_has_track_changes(output) is True
    import zipfile

    with zipfile.ZipFile(BytesIO(output)) as archive:
        assert "word/comments.xml" in archive.namelist()
        comments_xml = archive.read("word/comments.xml").decode()
        document_xml = archive.read("word/document.xml").decode()
    assert comments_xml.count("RULE: FINRA Rule 3110 requires") == 1
    assert comments_xml.count("RULE: FINRA Rule 3110(b) requires") == 1
    assert comments_xml.count("GAP: The policy omits") == 1
    assert comments_xml.count("GAP: The policy does not require") == 1
    assert document_xml.count("commentRangeStart") == 2
    assert document_xml.count("commentReference") == 2


def test_build_metadata_field_overlays_for_internal_policy_bumps_fields(monkeypatch) -> None:
    from app.services.tracked_redline_document import build_metadata_field_overlays_for_internal_policy

    monkeypatch.setattr(
        "app.services.policy_revision_document.policy_template_dates",
        lambda reference=None: {
            "effective_date": "July 14, 2026",
            "next_review_date": "July 14, 2028",
            "approved_date": "July 14, 2026",
        },
    )
    internal = aml_internal_docx()
    overlays = build_metadata_field_overlays_for_internal_policy(
        internal,
        applicable_regulation="FINRA Rule 3310(a); 31 CFR § 1023.320; FINRA Rule 3120",
    )
    by_label = {
        overlay.original_text.split(":", 1)[0]: overlay.revised_text
        for overlay in overlays
    }
    assert by_label["Version"] == "Version: 2.3"
    assert by_label["Effective Date"] == "Effective Date: July 14, 2026"
    assert "FINRA Rule 3310" in by_label["Applicable Regulation"]
    assert "31 CFR § 1023.320" in by_label["Applicable Regulation"]
    assert "FINRA Rule 3120" in by_label["Applicable Regulation"]
    assert by_label["Next Review Date"] == "Next Review Date: July 14, 2028"
    assert by_label["Approved date"] == "Approved date: July 14, 2026"


def test_build_metadata_field_overlays_sets_approval_date_alias_to_today(monkeypatch) -> None:
    from app.services.tracked_redline_document import build_metadata_field_overlays_for_internal_policy

    monkeypatch.setattr(
        "app.services.policy_revision_document.policy_template_dates",
        lambda reference=None: {
            "effective_date": "July 18, 2026",
            "next_review_date": "July 18, 2028",
            "approved_date": "July 18, 2026",
        },
    )
    document = Document()
    document.add_paragraph("POLICY: Test Policy")
    document.add_paragraph("Approval Date: January 1, 2020")
    document.add_heading("1. Purpose", level=1)
    document.add_paragraph("Body text.")
    buffer = BytesIO()
    document.save(buffer)
    overlays = build_metadata_field_overlays_for_internal_policy(buffer.getvalue())
    assert len(overlays) == 1
    assert overlays[0].original_text.startswith("Approval Date:")
    assert overlays[0].revised_text == "Approval Date: July 18, 2026"

    from app.services.tracked_redline_document import build_metadata_field_overlays_for_internal_policy

    monkeypatch.setattr(
        "app.services.policy_revision_document.policy_template_dates",
        lambda reference=None: {
            "effective_date": "July 17, 2026",
            "next_review_date": "July 17, 2028",
            "approved_date": "July 17, 2026",
        },
    )
    internal = aml_internal_docx()
    prose = (
        "FINRA Rules 3310; 2090; Bank Secrecy Act; 31 CFR Chapter X; The Purpose section must "
        "acknowledge the Firm's AI-driven analytical process for AML monitoring of overseas "
        "accounts. The Scope section must explicitly include model validation and data science "
        "oversight. The SAR filing workflow must clarify escalation under SEC Rule 17a-4. "
        "FINRA FINRA Regulations"
    )
    overlays = build_metadata_field_overlays_for_internal_policy(
        internal,
        applicable_regulation=prose,
    )
    applicable = next(
        overlay.revised_text
        for overlay in overlays
        if overlay.original_text.lower().startswith("applicable regulation:")
    )
    assert "Purpose section must" not in applicable
    assert "model validation" not in applicable
    assert "FINRA FINRA Regulations" not in applicable
    assert "FINRA Rule 3310" in applicable
    assert len(applicable) < 280


def aml_internal_docx() -> bytes:
    """POLICY-first template with a leading blank paragraph, like many AML uploads."""
    document = Document()
    document.add_paragraph("")
    policy_line = document.add_paragraph()
    policy_line.add_run("POLICY: Anti-Money Laundering (AML) Policy").bold = True
    for label, value in [
        ("COMPANY", "Acme Securities"),
        ("DEPARTMENT", "Compliance"),
        ("Version", "2.2"),
        ("Effective Date", "October 1, 2022"),
        ("Applicable Regulation", "FINRA Rule 3310"),
        ("Next Review Date", "October 1, 2024"),
        ("Approved date", "October 1, 2022"),
    ]:
        row = document.add_paragraph()
        row.add_run(f"{label}: ").bold = True
        row.add_run(value)
    document.add_heading("1. Purpose", level=1)
    document.add_paragraph("This policy establishes the AML program for the Firm.")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_apply_tracked_redline_overlay_skips_metadata_replacements() -> None:
    from docx_revisions import RevisionParagraph

    internal = aml_internal_docx()
    bad_next_review = (
        "Next Review Date: [INSERT CURRENT UPDATED DATE — conduct overdue review immediately "
        "and establish annual review cadence aligned with the Firm's FINRA Rule 3120 testing calendar]"
    )
    output = apply_tracked_redline_overlay(
        internal,
        replacements=[
            TrackedReplacement(
                search_text="Next Review Date: October 1, 2024",
                replace_text=bad_next_review,
                comment="RULE: x\n\nGAP: y\n\nFIX: z",
            ),
            TrackedReplacement(
                search_text="This policy establishes the AML program for the Firm.",
                replace_text="This policy establishes the Firm's written AML program.",
                comment="RULE: FINRA 3310\n\nGAP: Missing written program language.\n\nFIX: Clarifies written program.",
            ),
        ],
    )
    document = Document(BytesIO(output))
    accepted = [
        RevisionParagraph.from_paragraph(paragraph).accepted_text
        for paragraph in document.paragraphs
    ]
    assert any("Next Review Date: October 1, 2024" in text for text in accepted)
    assert not any("INSERT CURRENT UPDATED DATE" in text for text in accepted)
    assert any("written AML program" in text for text in accepted)


def test_extract_accepted_redline_text_includes_tracked_replacements() -> None:
    from app.services.redline_document import extract_policy_paragraphs, policy_text_from_paragraphs
    from app.services.tracked_redline_document import extract_accepted_redline_text

    internal = aml_internal_docx()
    output = apply_tracked_redline_overlay(
        internal,
        replacements=[
            TrackedReplacement(
                search_text="This policy establishes the AML program for the Firm.",
                replace_text="This policy establishes the Firm's written AML program under FINRA Rule 3310.",
                comment="RULE: FINRA Rule 3310\n\nGAP: Missing written program.\n\nFIX: Clarifies written program.",
            ),
        ],
    )
    accepted = extract_accepted_redline_text(output)
    plain = policy_text_from_paragraphs(extract_policy_paragraphs(output))
    assert "written AML program under FINRA Rule 3310" in accepted
    # Plain paragraph.text often drops tracked insertions — that path must not be used for similarity.
    assert "written AML program under FINRA Rule 3310" not in plain or "AML program for the Firm" in plain

    from app.services.tracked_redline_document import is_cover_or_metadata_line

    assert is_cover_or_metadata_line(
        "Next Review Date: [INSERT CURRENT UPDATED DATE — conduct overdue review]"
    )
    assert is_cover_or_metadata_line("Version: 2.2")
    assert not is_cover_or_metadata_line("This policy establishes the AML program for the Firm.")


def test_apply_paragraph_overlays_never_inserts_gap_summary_fallback() -> None:
    internal = aml_internal_docx()
    gap_summary = (
        "Based on the gap analysis between the existing AML policy and FINRA Rule 3310, "
        "several updates are required to ensure full compliance."
    )
    output = apply_paragraph_overlays(
        internal,
        [
            ParagraphOverlay(
                original_text="Paragraph that does not exist in the document.",
                revised_text="Replacement text.",
            )
        ],
        gap_summary=gap_summary,
    )
    import zipfile

    with zipfile.ZipFile(BytesIO(output)) as archive:
        xml = archive.read("word/document.xml").decode().lower()
    assert "based on the gap analysis" not in xml
