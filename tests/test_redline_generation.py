import json
from contextlib import contextmanager
from datetime import UTC, date, datetime
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from ai_fakes import FakeLLM
from docx import Document
from fastapi import HTTPException

from app.core.config import Settings
from app.core.security import CurrentUser
from app.models.core import DocumentChunk, GapAnalysis, Policy, PolicyRevision
from app.services.llm_json import parse_llm_json_dict
from app.services.policy_revision_document import (
    align_policy_sections_to_template,
    policy_template_dates,
)
from app.services.redline_generation_service import (
    REDLINE_GENERATION_SYSTEM_PROMPT,
    PolicyApprovalFooter,
    PolicyDocumentHeader,
    PolicyDocumentSectionRevision,
    PolicySectionRevision,
    RedlineContext,
    RedlineDraft,
    build_tracked_redline_docx,
    draft_replacement_language,
    enrich_redline_draft_from_findings,
    finalize_policy_redline_for_revision,
    generate_policy_redline,
    generate_policy_redline_internal,
    get_or_create_pending_revision,
    is_docx_bytes,
    load_regulatory_reference_sections_for_redline,
    parse_redline_fields_loosely,
    persist_revision_redline_url,
    validate_redline_draft,
)
from app.services.redline_vector_validation_service import RedlineVectorValidationResult
from app.services.storage_service import StoredFile

TEMPLATE_REFERENCE_DATE = date(2026, 6, 29)


@pytest.fixture(autouse=True)
def mock_redline_regulator_scrape(monkeypatch: pytest.MonkeyPatch) -> None:
    async def fake_scrape(*_args: object, **_kwargs: object) -> int:
        return 0

    monkeypatch.setattr(
        "app.services.internal_policy_workflow_service.scrape_and_refresh_all_regulator_references",
        fake_scrape,
    )
    monkeypatch.setattr(
        "app.services.internal_policy_workflow_service.refresh_org_government_reference_embeddings",
        fake_scrape,
    )


@pytest.fixture(autouse=True)
def mock_redline_policy_vector_loading(monkeypatch: pytest.MonkeyPatch) -> None:
    async def fake_load_policy_vectors(*_args: object, **_kwargs: object) -> list[list[float]]:
        return [[1.0, 0.0, 0.0]]

    monkeypatch.setattr(
        "app.services.redline_generation_service.load_policy_vectors_for_validation",
        fake_load_policy_vectors,
    )


@pytest.fixture(autouse=True)
def mock_redline_vector_comparisons(monkeypatch: pytest.MonkeyPatch) -> None:
    async def fake_persist_and_compare(*_args: object, **_kwargs: object) -> tuple[list[object], list[object]]:
        return [], []

    monkeypatch.setattr(
        "app.services.redline_generation_service._persist_and_compare_redline_vectors",
        fake_persist_and_compare,
    )


@pytest.fixture(autouse=True)
def mock_redline_revision_persist(monkeypatch: pytest.MonkeyPatch) -> None:
    def fake_persist(
        revision_id: str,
        org_id: str,
        redline_file_url: str,
        db: object | None = None,
        **_: object,
    ) -> PolicyRevision:
        if db is not None and hasattr(db, "added"):
            for item in db.added:
                if isinstance(item, PolicyRevision) and item.id == revision_id:
                    item.redline_file_url = redline_file_url
                    if hasattr(db, "committed"):
                        db.committed = True
                    return item
        return PolicyRevision(
            id=revision_id,
            gap_analysis_id="gap-id",
            policy_id="policy-id",
            org_id=org_id,
            redline_file_url=redline_file_url,
            status="PENDING",
            created_at=datetime.now(UTC),
        )

    monkeypatch.setattr(
        "app.services.redline_generation_service.persist_revision_redline_url",
        fake_persist,
    )


def sample_policy_docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Electronic Communications Policy", level=1)
    document.add_paragraph(
        "Supervisors review retained electronic communications every month and document exceptions."
    )
    document.add_paragraph(
        "Employees may not use personal devices for firm business unless approved in writing."
    )
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def sample_gap_analysis() -> GapAnalysis:
    return GapAnalysis(
        id="gap-id",
        org_id="org_example",
        policy_id="policy-id",
        trigger_type="publication",
        trigger_id="publication-id",
        impact_score="HIGH",
        gap_summary="FINRA requires documented approval evidence for covered communications.",
        regulatory_citations=[
            {
                "source": "FINRA",
                "reference": "https://www.finra.org/rules-guidance/notices/example",
                "quote": "document supervisory approval",
                "rationale": "The trigger adds approval evidence requirements.",
            }
        ],
    )


def sample_policy() -> Policy:
    return Policy(
        id="policy-id",
        org_id="org_example",
        name="Electronic Communications Policy",
        version=3,
        status="ACTIVE",
        file_url="storage://policy-documents/org_example/policies/policy-id/source.docx",
    )


def sample_context() -> RedlineContext:
    policy_bytes = sample_policy_docx_bytes()
    return RedlineContext(
        gap_analysis=sample_gap_analysis(),
        policy=sample_policy(),
        policy_bytes=policy_bytes,
        policy_text=(
            "Supervisors review retained electronic communications every month and document exceptions.\n"
            "Employees may not use personal devices for firm business unless approved in writing."
        ),
        regulatory_references=[
            {
                "policy_name": "FINRA Media Center Reference",
                "regulator": "FINRA",
                "excerpt": "Firms must document supervisory approval for covered communications.",
            }
        ],
        organization_name="Meridian Financial Services LLC",
    )


def sample_template_draft() -> RedlineDraft:
    dates = policy_template_dates(TEMPLATE_REFERENCE_DATE)
    return RedlineDraft(
        document_title="Electronic Communications Policy",
        drafting_summary="Adds approval evidence requirements.",
        regulatory_sources_used=[],
        policy_header=PolicyDocumentHeader(
            company_name="Meridian Financial Services LLC",
            company_department="Compliance & Legal Department",
            policy_name="Electronic Communications Policy",
            policy_number="CMP-EC-001",
            version="3",
            effective_date=dates["effective_date"],
            owner="Chief Compliance Officer / Chief Operating Officer",
            applicable_regulation="FINRA Rule 3110",
            next_review_date=dates["next_review_date"],
        ),
        policy_sections=[
            PolicyDocumentSectionRevision(
                number="1",
                title="Purpose",
                body=(
                    "Supervisors review retained electronic communications every month, "
                    "document exceptions, and retain supervisory approval evidence."
                ),
                subsections=[],
            )
        ],
        approval=PolicyApprovalFooter(
            approved_by="Sarah Mitchell, Chief Compliance Officer",
            approved_date=dates["approved_date"],
        ),
    )


def test_request_review_revision_only_marks_requested_gap() -> None:
    from app.services.redline_generation_service import request_review_revision

    gap_a = sample_gap_analysis()
    gap_b = sample_gap_analysis()
    gap_b.id = "gap-b"
    policy = sample_policy()
    context_a = RedlineContext(
        gap_analysis=gap_a,
        policy=policy,
        policy_bytes=b"policy",
        policy_text="Policy text",
        regulatory_references=[],
        organization_name="Example Org",
        approval_chain=None,
    )
    revisions = {
        "rev-a": PolicyRevision(
            id="rev-a",
            gap_analysis_id=gap_a.id,
            policy_id=policy.id,
            org_id=policy.org_id,
            redline_file_url=None,
            status="PENDING",
            created_at=datetime.now(UTC),
        ),
        "rev-b": PolicyRevision(
            id="rev-b",
            gap_analysis_id=gap_b.id,
            policy_id=policy.id,
            org_id=policy.org_id,
            redline_file_url=None,
            status="PENDING",
            created_at=datetime.now(UTC),
        ),
    }

    class Query:
        def __init__(self, session: "SessionStub", mode: str = "select") -> None:
            self.session = session
            self.mode = mode
            self.filters: dict[str, object | None] = {}

        def filter(self, *_args: object) -> "Query":
            for arg in _args:
                if not hasattr(arg, "left"):
                    continue
                key = getattr(arg.left, "key", None) or getattr(arg.left, "name", None)
                if key is None:
                    continue
                if hasattr(arg, "right") and hasattr(arg.right, "value"):
                    self.filters[str(key)] = arg.right.value
                else:
                    self.filters[str(key)] = None
            return self

        def order_by(self, *_args: object) -> "Query":
            return self

        def first(self) -> PolicyRevision | None:
            for revision in revisions.values():
                if not all(
                    getattr(revision, key, None) == value
                    for key, value in self.filters.items()
                ):
                    continue
                return revision
            return None

        def update(self, values: dict[object, object], synchronize_session: bool = False) -> int:
            del synchronize_session
            revision_id = self.filters.get("id")
            revision = revisions.get(str(revision_id))
            if revision is None:
                return 0
            for key, expected in self.filters.items():
                if key == "id":
                    continue
                if getattr(revision, key, None) != expected:
                    return 0
            requested_at = values.get(PolicyRevision.review_requested_at)
            revision.review_requested_at = requested_at  # type: ignore[assignment]
            return 1

    class SessionStub:
        def query(self, *_entities: object) -> Query:
            return Query(self)

        def add(self, _value: object) -> None:
            return None

        def commit(self) -> None:
            return None

        def refresh(self, _value: object) -> None:
            return None

        def rollback(self) -> None:
            return None

    session = SessionStub()
    activated, newly_requested = request_review_revision(context_a, session)  # type: ignore[arg-type]

    assert activated.id == "rev-a"
    assert newly_requested is True
    assert revisions["rev-a"].review_requested_at is not None
    assert revisions["rev-b"].review_requested_at is None


def test_get_or_create_pending_revision_reuses_open_pending() -> None:
    gap_analysis = sample_gap_analysis()
    policy = sample_policy()
    pending = PolicyRevision(
        id="rev-pending",
        gap_analysis_id=gap_analysis.id,
        policy_id=policy.id,
        org_id=policy.org_id,
        redline_file_url=None,
        status="PENDING",
        created_at=datetime.now(UTC),
    )

    class Query:
        def __init__(self, result: object) -> None:
            self.result = result

        def filter(self, *_args: object) -> object:
            return self

        def order_by(self, *_args: object) -> object:
            return self

        def first(self) -> object:
            return self.result

        def update(self, values: dict[object, object], synchronize_session: bool = False) -> int:
            del synchronize_session, values
            if isinstance(self.result, PolicyRevision) and self.result.review_requested_at is None:
                self.result.review_requested_at = datetime.now(UTC)
                return 1
            return 0

    class SessionStub:
        def __init__(self, result: object) -> None:
            self.result = result
            self.added: list[object] = []

        def query(self, *_entities: object) -> Query:
            return Query(self.result)

        def add(self, value: object) -> None:
            self.added.append(value)

        def commit(self) -> None:
            return None

        def refresh(self, _value: object) -> None:
            return None

        def rollback(self) -> None:
            return None

    context = sample_context()
    reused = get_or_create_pending_revision(context, SessionStub(pending))  # type: ignore[arg-type]
    assert reused.id == "rev-pending"
    assert reused.status == "PENDING"

    created = get_or_create_pending_revision(context, SessionStub(None))  # type: ignore[arg-type]
    assert created.status == "PENDING"
    assert created.gap_analysis_id == gap_analysis.id
    assert created.redline_file_url is None


@pytest.mark.asyncio
async def test_generate_policy_redline_internal_can_defer_generation(monkeypatch) -> None:
    gap_analysis = sample_gap_analysis()
    policy = sample_policy()
    pending = PolicyRevision(
        id="rev-pending",
        gap_analysis_id=gap_analysis.id,
        policy_id=policy.id,
        org_id=policy.org_id,
        redline_file_url=None,
        status="PENDING",
        created_at=datetime.now(UTC),
    )

    class Query:
        def __init__(self, result: object) -> None:
            self.result = result

        def join(self, *_args: object) -> object:
            return self

        def filter(self, *_args: object) -> object:
            return self

        def order_by(self, *_args: object) -> object:
            return self

        def one_or_none(self) -> object:
            return self.result

        def first(self) -> object:
            return self.result

        def all(self) -> object:
            return []

        def update(self, *_args: object, **_kwargs: object) -> int:
            return 1

    class SessionStub:
        def query(self, *entities: object) -> Query:
            if entities == (GapAnalysis, Policy):
                return Query((gap_analysis, policy))
            if entities == (PolicyRevision,):
                return Query(pending)
            if entities == (Policy,):
                return Query([])
            if entities == (DocumentChunk,):
                return Query([])
            return Query(None)

        def add(self, _value: object) -> None:
            return None

        def commit(self) -> None:
            return None

        def refresh(self, _value: object) -> None:
            return None

        def rollback(self) -> None:
            return None

    monkeypatch.setattr(
        "app.services.redline_generation_service.download_file",
        lambda _file_url: sample_policy_docx_bytes(),
    )
    monkeypatch.setattr(
        "app.services.redline_generation_service.load_regulatory_reference_sections_for_redline",
        lambda *_args, **_kwargs: sample_context().regulatory_references,
    )
    finalize = AsyncMock()
    notify = AsyncMock()
    monkeypatch.setattr(
        "app.services.redline_generation_service.finalize_policy_redline_for_revision",
        finalize,
    )
    monkeypatch.setattr(
        "app.services.redline_generation_service.notify_new_revision_from_payload",
        notify,
    )

    response = await generate_policy_redline_internal(
        "gap-id",
        "org_example",
        SessionStub(),  # type: ignore[arg-type]
        defer_generation=True,
    )

    assert response.status == "PENDING"
    assert response.redline_file_url is None
    finalize.assert_not_awaited()
    notify.assert_awaited_once()


def test_schedule_pending_redline_generations_queues_missing_files() -> None:
    scheduled: list[tuple[str, str, str]] = []

    class BackgroundTasksStub:
        def add_task(self, func: object, *args: object) -> None:
            scheduled.append((str(args[0]), str(args[1]), str(args[2])))

    from app.services.redline_generation_service import schedule_pending_redline_generations

    schedule_pending_redline_generations(
        BackgroundTasksStub(),
        "org_example",
        [
            {
                "revision_id": "rev-1",
                "gap_analysis_id": "gap-1",
                "redline_url": None,
            },
            {
                "revision_id": "rev-2",
                "gap_analysis_id": "gap-2",
                "redline_url": "storage://policy-redlines/org_example/rev-2/redline.docx",
            },
        ],
    )

    assert scheduled == [("rev-1", "gap-1", "org_example")]


@pytest.mark.asyncio
async def test_finalize_policy_redline_for_revision_uses_fallback_on_drafting_failure(monkeypatch) -> None:
    context = sample_context()
    revision = PolicyRevision(
        id="rev-pending",
        gap_analysis_id=context.gap_analysis.id,
        policy_id=context.policy.id,
        org_id=context.policy.org_id,
        redline_file_url=None,
        status="PENDING",
        created_at=datetime.now(UTC),
    )

    class SessionStub:
        def commit(self) -> None:
            return None

        def refresh(self, _value: object) -> None:
            return None

        def rollback(self) -> None:
            return None

    async def failing_draft(*_args: object, **_kwargs: object) -> RedlineDraft:
        raise HTTPException(status_code=502, detail="Claude redline drafting failed")

    monkeypatch.setattr(
        "app.services.redline_generation_service.draft_replacement_language",
        failing_draft,
    )
    monkeypatch.setattr(
        "app.services.redline_generation_service.upload_file",
        lambda *_args, **_kwargs: StoredFile(
            bucket="policy-redlines",
            path="org/policy-revisions/rev-pending/file.docx",
        ),
    )
    monkeypatch.setattr(
        "app.services.redline_generation_service.notify_new_revision",
        AsyncMock(),
    )
    monkeypatch.setattr(
        "app.services.redline_generation_service.validate_redline_against_regulator_documents",
        AsyncMock(
            return_value=RedlineVectorValidationResult(
                accepted=True,
                similarity=1.0,
                threshold=0.25,
            )
        ),
    )
    monkeypatch.setattr(
        "app.services.redline_generation_service.load_all_regulator_reference_embeddings",
        lambda *_args, **_kwargs: [[1.0, 0.0, 0.0]],
    )

    response = await finalize_policy_redline_for_revision(
        revision,
        context,
        SessionStub(),  # type: ignore[arg-type]
    )

    assert response.redline_file_url is not None


def test_load_regulatory_reference_sections_for_redline_prefers_gap_excerpts() -> None:
    gap_analysis = sample_gap_analysis()
    gap_analysis.retrieval_metadata = {
        "government_reference_excerpts": [
            {
                "reference_policy": "FINRA Media Center Reference",
                "regulator_source": "FINRA",
                "chunk_index": 0,
                "text": "Firms must document supervisory approval for covered communications.",
            }
        ]
    }

    class Query:
        def __init__(self, result: object) -> None:
            self.result = result

        def filter(self, *_args: object) -> object:
            return self

        def order_by(self, *_args: object) -> object:
            return self

        def limit(self, *_args: object) -> object:
            return self

        def all(self) -> object:
            return self.result

        def first(self) -> object:
            return self.result

    class SessionStub:
        def query(self, *_entities: object) -> Query:
            return Query([])

    references = load_regulatory_reference_sections_for_redline(
        SessionStub(),  # type: ignore[arg-type]
        "org_example",
        sample_policy(),
        gap_analysis,
        sample_context().policy_text,
    )

    assert len(references) == 1
    assert references[0]["policy_name"] == "FINRA Media Center Reference"
    assert references[0]["match_source"] == "gap_analysis_retrieval"


def test_validate_redline_draft_keeps_valid_sections() -> None:
    context = sample_context()
    draft = validate_redline_draft(
        RedlineDraft(
            document_title="Revision Draft",
            sections=[
                PolicySectionRevision(
                    heading="Supervision",
                    original_text="document exceptions.",
                    revised_text="document exceptions and retain approval evidence.",
                    change_note="Adds approval evidence.",
                )
            ],
            drafting_summary="Adds approval evidence.",
        ),
        context,
    )

    assert len(draft.policy_sections) == 7
    assert draft.policy_sections[0].title == "Purpose"
    assert "retain approval evidence" in draft.policy_sections[0].body


def test_parse_redline_json_recovers_sections_array() -> None:
    text = """{
  "document_title": "Revision Draft",
  "drafting_summary": "Adds approval evidence.",
  "sections": [
    {
      "heading": "Supervision",
      "original_text": "document exceptions.",
      "revised_text": "document exceptions and retain approval evidence.",
      "change_note": "FINRA approval citation."
    }
  ]
}"""
    payload = parse_llm_json_dict(text, loose_parser=parse_redline_fields_loosely)
    draft = RedlineDraft.model_validate(payload)
    assert len(draft.sections) == 1
    assert draft.sections[0].revised_text.endswith("approval evidence.")


@pytest.mark.asyncio
async def test_opus_revision_prompt_uses_matched_regulatory_reference_material() -> None:
    llm = FakeLLM(
        json.dumps(
            {
                "document_title": "Revision Draft",
                "drafting_summary": "Adds approval evidence.",
                "edits": [
                    {
                        "anchor_text": "document exceptions.",
                        "replacement_text": (
                            "document exceptions and retain documented supervisory approval evidence."
                        ),
                        "citation_comment": (
                            "RULE: FINRA Rule 3110 — supervisory procedures.\\n\\n"
                            "GAP: Exceptions omit documented approval evidence.\\n\\n"
                            "FIX: Retain written supervisory approval for each exception."
                        ),
                    }
                ],
            }
        )
    )

    draft = await draft_replacement_language(
        sample_context(),
        Settings(ANTHROPIC_API_KEY="test-key"),
        llm,
    )

    assert len(draft.edits) == 1
    assert "retain documented supervisory approval evidence" in draft.edits[0].replacement_text
    assert "RULE:" in draft.edits[0].citation_comment
    request = llm.calls[0]
    assert request["system"] == REDLINE_GENERATION_SYSTEM_PROMPT
    assert request["cache_system"] is True
    assert request["throttle"] == "sonnet"
    assert request["temperature"] == 0
    user_payload = json.loads(str(request["messages"][0].content))
    assert user_payload["task"] == "sentence_level_internal_policy_redline_edits"
    assert user_payload["output_format"] == "sentence_edits_with_unique_rule_gap_fix_comments"
    assert "organization_context" in user_payload
    assert "matched_regulatory_reference_material" in user_payload
    assert "gap_analysis_context" in user_payload
    assert "sentence_findings" in user_payload["gap_analysis_context"]
    assert user_payload["internal_policy"]["paragraphs"][0]["text"].startswith("Supervisors review")


@pytest.mark.asyncio
async def test_generate_policy_redline_uploads_docx_and_persists_pending_revision(monkeypatch) -> None:
    gap_analysis = sample_gap_analysis()
    policy = sample_policy()
    organization = SimpleNamespace(
        auth_org_id="org_example",
        name="Meridian Financial Services LLC",
        approval_chain={
            "assigned_reviewer_email": "reviewer@example.com",
            "assigned_reviewer_name": "Reviewer",
        },
    )
    uploaded: dict[str, object] = {}

    class Query:
        def __init__(self, result: object) -> None:
            self.result = result

        def join(self, *_args: object) -> object:
            return self

        def filter(self, *_args: object) -> object:
            return self

        def order_by(self, *_args: object) -> object:
            return self

        def one_or_none(self) -> object:
            return self.result

        def first(self) -> object:
            return self.result

        def all(self) -> object:
            return self.result if isinstance(self.result, list) else []

    class RecordingSession:
        def __init__(self) -> None:
            self.added: list[object] = []
            self.committed = False

        def query(self, *entities: object) -> Query:
            if entities == (GapAnalysis, Policy):
                return Query((gap_analysis, policy))
            if entities == (PolicyRevision,):
                return Query(None)
            if entities == (Policy,):
                return Query([])
            if entities == (DocumentChunk,):
                return Query([])
            if entities and getattr(entities[0], "__name__", "") == "Organization":
                return Query(organization)
            return Query(None)

        def add(self, value: object) -> None:
            self.added.append(value)

        def commit(self) -> None:
            self.committed = True

        def refresh(self, _value: object) -> None:
            return None

        def rollback(self) -> None:
            return None

    def fake_upload_file(bucket: str, path: str, contents: bytes, content_type: str | None) -> StoredFile:
        uploaded.update({"bucket": bucket, "path": path, "contents": contents, "content_type": content_type})
        return StoredFile(bucket=bucket, path=path)

    monkeypatch.setattr(
        "app.services.redline_generation_service.download_file",
        lambda _file_url: sample_policy_docx_bytes(),
    )
    monkeypatch.setattr("app.services.redline_generation_service.upload_file", fake_upload_file)
    monkeypatch.setattr(
        "app.services.redline_generation_service.load_regulatory_reference_sections_for_redline",
        lambda *_args, **_kwargs: sample_context().regulatory_references,
    )
    mock_notify = AsyncMock(return_value={"status": "sent"})
    monkeypatch.setattr(
        "app.services.redline_generation_service.send_revision_notification",
        mock_notify,
    )
    monkeypatch.setattr(
        "app.services.redline_generation_service.validate_redline_against_regulator_documents",
        AsyncMock(
            return_value=RedlineVectorValidationResult(
                accepted=True,
                similarity=0.42,
                threshold=0.25,
                redline_vectors=[[1.0, 0.0, 0.0]],
            )
        ),
    )
    monkeypatch.setattr(
        "app.services.redline_generation_service.load_all_regulator_reference_embeddings",
        lambda *_args, **_kwargs: [[1.0, 0.0, 0.0]],
    )

    db = RecordingSession()
    llm = FakeLLM(
        json.dumps(
            {
                "drafting_summary": "Adds approval evidence.",
                "edits": [
                    {
                        "anchor_text": (
                            "Supervisors review retained electronic communications every month "
                            "and document exceptions."
                        ),
                        "replacement_text": (
                            "Supervisors review retained electronic communications every month "
                            "and document exceptions and retain approval evidence."
                        ),
                        "citation_comment": (
                            "RULE: FINRA Rule 3110\n\n"
                            "GAP: The policy omits approval evidence.\n\n"
                            "FIX: Retain supervisory approval evidence."
                        ),
                    }
                ],
            }
        )
    )
    response = await generate_policy_redline(
        "gap-id",
        CurrentUser(user_id="cco-id", org_id="org_example", role="chief_compliance_officer"),
        db,  # type: ignore[arg-type]
        client=llm,
    )

    persisted = [item for item in db.added if isinstance(item, PolicyRevision)]
    assert len(persisted) >= 1
    assert persisted[0].status == "PENDING"
    assert db.committed is True
    assert response.status == "PENDING"
    assert isinstance(uploaded["contents"], bytes)
    assert is_docx_bytes(uploaded["contents"]) is True
    mock_notify.assert_awaited_once()


def test_build_tracked_redline_docx_applies_policy_section_body_changes() -> None:
    output = build_tracked_redline_docx(
        sample_policy_docx_bytes(),
        sample_context().policy_text,
        RedlineDraft(
            policy_sections=[
                PolicyDocumentSectionRevision(
                    number="1",
                    title="Electronic Communications Policy",
                    body=(
                        "Supervisors review retained electronic communications every month "
                        "and document exceptions and retain approval evidence."
                    ),
                    subsections=[],
                )
            ],
            drafting_summary="Adds approval evidence.",
        ),
    )
    assert is_docx_bytes(output) is True
    from app.services.tracked_redline_document import (
        accept_all_tracked_changes,
        docx_has_track_changes,
    )

    assert docx_has_track_changes(output) is True
    document = Document(BytesIO(accept_all_tracked_changes(output)))
    assert any("approval evidence" in paragraph.text for paragraph in document.paragraphs)


def test_enrich_redline_draft_from_findings_uses_body_paragraph_when_sentence_missing() -> None:
    context = sample_context()
    context.gap_analysis.rule = "FINRA Rule 3110"
    context.gap_analysis.gap_summary = "The policy omits approval evidence."
    context.gap_analysis.fix = (
        "Supervisors review retained electronic communications every month "
        "and document exceptions and retain approval evidence."
    )
    draft = enrich_redline_draft_from_findings(
        RedlineDraft(drafting_summary="fallback"),
        context,
    )
    assert draft.edits
    assert "approval evidence" in draft.edits[0].replacement_text
    assert "Supervisors review" in draft.edits[0].anchor_text


def test_build_tracked_redline_docx_applies_legacy_section_replacements() -> None:
    output = build_tracked_redline_docx(
        sample_policy_docx_bytes(),
        sample_context().policy_text,
        RedlineDraft(
            sections=[
                PolicySectionRevision(
                    original_text="document exceptions.",
                    revised_text="document exceptions and retain approval evidence.",
                )
            ],
            drafting_summary="Adds approval evidence.",
        ),
    )
    assert is_docx_bytes(output) is True
    from app.services.tracked_redline_document import docx_has_track_changes

    assert docx_has_track_changes(output) is True


def test_regulations_mentioned_in_redline_extracts_citations_not_prose() -> None:
    from app.services.redline_generation_service import (
        RedlineEdit,
        RegulatorySourceNote,
        regulations_mentioned_in_redline,
    )

    draft = RedlineDraft(
        edits=[
            RedlineEdit(
                anchor_text="Purpose.",
                replacement_text="Purpose including AI AML monitoring.",
                citation_comment=(
                    "RULE: FINRA Rules 3310 and Bank Secrecy Act. The Purpose section must "
                    "acknowledge the Firm's AI-driven analytical process for AML monitoring of "
                    "overseas accounts and require model validation oversight.\n\n"
                    "GAP: Policy omits AI monitoring controls.\n\n"
                    "FIX: Add AI AML scope language."
                ),
            ),
            RedlineEdit(
                anchor_text="Scope.",
                replacement_text="Scope including overseas accounts.",
                citation_comment=(
                    "RULE: SEC Rule 17a-4 and FINRA Rule 3120 require retention and testing.\n\n"
                    "GAP: Scope omits overseas accounts.\n\n"
                    "FIX: Expand scope."
                ),
            ),
        ],
        regulatory_sources_used=[
            RegulatorySourceNote(
                policy_name="FINRA Regulations",
                regulator="FINRA",
                excerpt="Long excerpt that should not be copied into Applicable Regulation " * 5,
            )
        ],
    )
    result = regulations_mentioned_in_redline(draft)
    assert result == "FINRA Rule 3310; SEC Rule 17a-4; FINRA Rule 3120"
    assert "Purpose section must" not in result
    assert "Bank Secrecy Act" not in result
    assert "model validation" not in result


def test_bump_policy_version_label_always_increments_one() -> None:
    from app.services.policy_revision_document import bump_policy_version_label

    assert bump_policy_version_label("1.2") == "1.3"
    assert bump_policy_version_label("v1.2") == "v1.3"
    assert bump_policy_version_label("V2.9") == "V2.10"
    assert bump_policy_version_label("Version 3.0") == "3.1"
    assert bump_policy_version_label("4") == "5"

def test_policy_template_dates_use_today_and_two_year_review() -> None:
    dates = policy_template_dates(TEMPLATE_REFERENCE_DATE)
    assert dates["effective_date"] == "June 29, 2026"
    assert dates["next_review_date"] == "June 29, 2028"
    assert dates["approved_date"] == "June 29, 2026"


def test_policy_paragraphs_for_prompt_omits_cover_metadata() -> None:
    from app.services.redline_generation_service import policy_paragraphs_for_prompt
    from tests.test_tracked_redline_document import aml_internal_docx

    paragraphs = policy_paragraphs_for_prompt(aml_internal_docx())
    joined = "\n".join(item["text"] for item in paragraphs)
    assert "Next Review Date" not in joined
    assert "Version:" not in joined
    assert "Effective Date" not in joined
    assert "AML program" in joined


def test_build_tracked_redline_docx_sets_next_review_two_years_and_ignores_llm_metadata(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.services.redline_generation_service import RedlineEdit
    from tests.test_tracked_redline_document import aml_internal_docx

    monkeypatch.setattr(
        "app.services.policy_revision_document.policy_template_dates",
        lambda reference=None: {
            "effective_date": "July 18, 2026",
            "next_review_date": "July 18, 2028",
            "approved_date": "July 18, 2026",
        },
    )
    policy_bytes = aml_internal_docx()
    draft = RedlineDraft(
        document_title="AML",
        drafting_summary="test",
        edits=[
            RedlineEdit(
                anchor_text="Next Review Date: October 1, 2024",
                replacement_text=(
                    "Next Review Date: [INSERT CURRENT UPDATED DATE — conduct overdue review]"
                ),
                citation_comment="RULE: a\n\nGAP: b\n\nFIX: c",
            ),
            RedlineEdit(
                anchor_text="This policy establishes the AML program for the Firm.",
                replacement_text="This policy establishes the Firm's written AML program.",
                citation_comment="RULE: FINRA 3310\n\nGAP: Missing written program.\n\nFIX: Clarifies written program.",
            ),
        ],
    )
    output = build_tracked_redline_docx(policy_bytes, "unused", draft)
    from docx_revisions import RevisionParagraph

    document = Document(BytesIO(output))
    accepted = [
        RevisionParagraph.from_paragraph(paragraph).accepted_text
        for paragraph in document.paragraphs
    ]
    assert any(text.startswith("Next Review Date: July 18, 2028") for text in accepted)
    assert not any("INSERT CURRENT UPDATED DATE" in text for text in accepted)
    assert any("written AML program" in text for text in accepted)


def test_dedupe_redline_edits_drops_metadata_anchors() -> None:
    from app.services.redline_generation_service import RedlineEdit, dedupe_redline_edits

    cleaned = dedupe_redline_edits(
        [
            RedlineEdit(
                anchor_text="Next Review Date: October 1, 2024",
                replacement_text=(
                    "Next Review Date: [INSERT CURRENT UPDATED DATE — conduct overdue review]"
                ),
                citation_comment="RULE: a\n\nGAP: b\n\nFIX: c",
            ),
            RedlineEdit(
                anchor_text="This policy establishes the AML program for the Firm.",
                replacement_text="This policy establishes the Firm's written AML program.",
                citation_comment="RULE: FINRA 3310\n\nGAP: Missing written program.\n\nFIX: Clarifies written program.",
            ),
        ]
    )
    assert len(cleaned) == 1
    assert "AML program" in cleaned[0].anchor_text


def test_align_policy_sections_to_template_preserves_canonical_titles() -> None:
    aligned = align_policy_sections_to_template(
        [
            {
                "number": "1",
                "title": "Purpose",
                "body": "Defines the policy objective.",
                "subsections": [],
            }
        ]
    )

    assert [section["title"] for section in aligned] == [
        "Purpose",
        "Scope",
        "Vendor Onboarding and Due Diligence",
        "Ongoing Vendor Monitoring",
        "Vendor Termination",
        "Roles and Responsibilities",
        "Recordkeeping",
    ]
    assert aligned[0]["body"] == "Defines the policy objective."
    assert aligned[2]["subsections"][0]["number"] == "3.2"
    assert aligned[2]["subsections"][0]["title"] == "Contract Requirements"


def test_redline_system_prompt_requires_sentence_level_unique_comments() -> None:
    assert "Analyze EVERY sentence" in REDLINE_GENERATION_SYSTEM_PROMPT
    assert "Sentence-by-sentence redline drafting" in REDLINE_GENERATION_SYSTEM_PROMPT
    assert "holistic meaning" in REDLINE_GENERATION_SYSTEM_PROMPT
    assert "business_change" in REDLINE_GENERATION_SYSTEM_PROMPT
    assert "marketing_content" in REDLINE_GENERATION_SYSTEM_PROMPT
    assert "Do NOT change section titles" in REDLINE_GENERATION_SYSTEM_PROMPT
    assert "NEVER revise cover-page or policy metadata fields yourself" in REDLINE_GENERATION_SYSTEM_PROMPT
    assert "Applicable Regulation: only compact rule name + number labels" in REDLINE_GENERATION_SYSTEM_PROMPT
    assert "Version: bump one minor step" in REDLINE_GENERATION_SYSTEM_PROMPT
    assert "Effective Date: today's date" in REDLINE_GENERATION_SYSTEM_PROMPT
    assert "Approved date / Approval date" in REDLINE_GENERATION_SYSTEM_PROMPT
    assert "Next Review Date: two years from today's date" in REDLINE_GENERATION_SYSTEM_PROMPT
    assert "citation_comment" in REDLINE_GENERATION_SYSTEM_PROMPT
    assert "RULE:" in REDLINE_GENERATION_SYSTEM_PROMPT
    assert "GAP and FIX wording must be unique per edit" in REDLINE_GENERATION_SYSTEM_PROMPT
    assert "Do not return policy_sections" in REDLINE_GENERATION_SYSTEM_PROMPT


def test_validate_redline_draft_overrides_header_dates_with_real_values(monkeypatch) -> None:
    monkeypatch.setattr(
        "app.services.redline_generation_service.policy_template_dates",
        lambda reference=None: policy_template_dates(TEMPLATE_REFERENCE_DATE),
    )
    context = sample_context()
    draft = validate_redline_draft(
        RedlineDraft(
            document_title="Revision Draft",
            policy_header=PolicyDocumentHeader(
                effective_date="January 1, 2020",
                next_review_date="January 1, 2021",
            ),
            policy_sections=[
                PolicyDocumentSectionRevision(
                    number="1",
                    title="Purpose",
                    body="Defines the policy objective.",
                    subsections=[],
                )
            ],
            approval=PolicyApprovalFooter(
                approved_by="Sarah Mitchell, Chief Compliance Officer",
                approved_date="January 1, 2020",
            ),
            drafting_summary="Updates purpose language.",
        ),
        context,
    )

    assert draft.policy_header.effective_date == "June 29, 2026"
    assert draft.policy_header.next_review_date == "June 29, 2028"
    assert draft.approval.approved_date == "June 29, 2026"


def test_persist_revision_redline_url_keeps_url_through_commit(monkeypatch) -> None:
    revision = PolicyRevision(
        id="rev-1",
        gap_analysis_id="gap-1",
        policy_id="policy-1",
        org_id="org_example",
        redline_file_url=None,
        status="PENDING",
        created_at=datetime.now(UTC),
    )
    stored_url = "storage://policy-redlines/org_example/policy-revisions/rev-1/CFTC_v1_redline.docx"
    committed_url: str | None = None

    class SessionStub:
        def query(self, _model: object) -> object:
            return self

        def filter(self, *_args: object) -> object:
            return self

        def one_or_none(self) -> PolicyRevision:
            return revision

        def flush(self) -> None:
            return None

        def refresh(self, _value: PolicyRevision) -> None:
            revision.redline_file_url = None

        def commit(self) -> None:
            nonlocal committed_url
            committed_url = revision.redline_file_url

        def expunge(self, _value: PolicyRevision) -> None:
            return None

    @contextmanager
    def fake_session_scope(*, commit: bool = True):
        session = SessionStub()
        yield session
        if commit:
            session.commit()

    monkeypatch.setattr(
        "app.services.redline_generation_service.session_scope",
        fake_session_scope,
    )

    result = persist_revision_redline_url("rev-1", "org_example", stored_url, db=None)

    assert result.redline_file_url == stored_url
    assert committed_url == stored_url


def test_gap_analysis_eligible_for_redline() -> None:
    from app.services.redline_generation_service import gap_analysis_eligible_for_redline

    assert gap_analysis_eligible_for_redline(trigger_type='publication', impact_score='HIGH') is True
    assert gap_analysis_eligible_for_redline(trigger_type='policy_update', impact_score='MEDIUM') is True
    assert gap_analysis_eligible_for_redline(trigger_type='publication', impact_score='INFO') is False
    assert gap_analysis_eligible_for_redline(trigger_type='business_change', impact_score='HIGH') is False
    assert gap_analysis_eligible_for_redline(
        trigger_type='business_change',
        impact_score='HIGH',
        retrieval_metadata={'business_change_policy_match': True, 'redline_eligible': True},
    ) is True
    assert gap_analysis_eligible_for_redline(
        trigger_type='business_change',
        impact_score='HIGH',
        retrieval_metadata={'new_policy_from_business_change': True, 'redline_eligible': False},
    ) is False


def test_schedule_missing_gap_analysis_redlines_is_noop() -> None:
    from app.schemas.gap_analyses import GapAnalysisResponse
    from app.services.redline_generation_service import schedule_missing_gap_analysis_redlines

    queued: list[tuple[str, str]] = []

    class BackgroundTasksStub:
        def add_task(self, func: object, gap_id: str, org_id: str) -> None:
            queued.append((gap_id, org_id))

    analyses = [
        GapAnalysisResponse(
            id='gap-high',
            policy_id='policy-1',
            policy_name='FINRA WSP',
            policy_version=1,
            impact_score='HIGH',
            trigger_type='publication',
            trigger_id='pub-1',
            trigger_label='Publication',
            gap_summary='Summary',
            retrieval_mode='VECTOR',
            analysis_mode='ANTHROPIC',
            created_at=datetime.now(tz=UTC),
            redline_ready=False,
        ),
    ]

    schedule_missing_gap_analysis_redlines(BackgroundTasksStub(), 'org_example', analyses)

    assert queued == []


def test_citation_comment_prefers_draft_short_rule_gap_fix() -> None:
    from app.services.redline_generation_service import citation_comment

    gap = GapAnalysis(
        id='gap-1',
        org_id='org_example',
        policy_id='policy-1',
        trigger_type='policy_update',
        trigger_id='policy-1',
        impact_score='HIGH',
        gap_summary=' '.join(['Long gap sentence that exceeds comment readability limits.'] * 20),
        regulatory_citations=[],
        rule='FINRA Rule 3110',
        fix='Add quarterly branch inspection procedures with documented exception follow-up by the CCO.',
    )
    draft = RedlineDraft(
        drafting_summary='Adds supervision language.',
        citation_comment=(
            'RULE: FINRA Rule 3110.\n\n'
            'GAP: Policy omits monthly exception documentation.\n\n'
            'FIX: Require supervisors to retain approval evidence each month.'
        ),
    )
    comment = citation_comment(gap, draft)
    assert 'RULE: FINRA Rule 3110.' in comment
    assert 'GAP: Policy omits monthly exception documentation.' in comment
    assert 'FIX: Require supervisors to retain approval evidence each month.' in comment

    fallback = citation_comment(gap)
    assert fallback.startswith('RULE:')
    assert '\n\nGAP:' in fallback
    assert '\n\nFIX:' in fallback
    assert 'Long gap sentence that exceeds comment readability limits.' in fallback
    assert 'Add quarterly branch inspection procedures' in fallback
    assert '…' not in fallback


def test_should_auto_generate_redlines_always_false() -> None:
    from app.schemas.agent import AgentRunRequest
    from app.services.gap_analysis_service import _should_auto_generate_redlines

    assert _should_auto_generate_redlines(AgentRunRequest(trigger_type='publication', trigger_id='x')) is False
    assert _should_auto_generate_redlines(
        AgentRunRequest(trigger_type='publication', trigger_id='x', auto_generate_redlines=True)
    ) is False
