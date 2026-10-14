from collections.abc import Generator
from io import BytesIO
from uuid import uuid4

from docx import Document
from fastapi.testclient import TestClient

from app.core.config import Settings
from app.db.session import get_db
from app.main import app
from app.models.core import DocumentChunk, Policy
from app.schemas.context_documents import ContextDocumentType
from app.services.document_ingestion import TextChunk
from app.services.policy_service import persist_policy_upload
from app.services.storage_service import (
    StoredFile,
    parse_storage_uri,
    storage_path,
    storage_url,
)

DOCX_MIME_TYPE = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'


def document_bytes() -> bytes:
    document = Document()
    document.add_heading('Exam response', level=1)
    document.add_paragraph(
        'The firm remediated its supervisory exception and documented the review.')
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def override_db_none() -> Generator[None, None, None]:
    yield None


def test_context_document_upload_supports_prior_exam_responses_without_db() -> None:
    app.dependency_overrides[get_db] = override_db_none
    client = TestClient(app)

    response = client.post(
        '/api/v1/context-documents/PRIOR_EXAM_RESPONSE/upload',
        files={'file': ('exam-response.docx',
                        document_bytes(), DOCX_MIME_TYPE)},
        headers={'x-org-id': 'test-org'},
    )

    assert response.status_code == 202
    payload = response.json()
    assert payload['org_id'] == 'test-org'
    assert payload['document_type'] == ContextDocumentType.PRIOR_EXAM_RESPONSE.value
    assert payload['persisted'] is False
    assert payload['chunk_count'] == 1
    app.dependency_overrides.clear()


def test_approval_chain_input_is_available_in_local_mode() -> None:
    app.dependency_overrides[get_db] = override_db_none
    client = TestClient(app)

    response = client.put(
        '/api/v1/organizations/current/approval-chain',
        json={
            'org_structure': [
                {
                    'name': 'Taylor Morgan',
                    'title': 'Chief Compliance Officer',
                    'user_id': 'user_cco',
                    'department': 'Compliance',
                }
            ],
            'approving_authorities': [
                {
                    'name': 'Taylor Morgan',
                    'title': 'Chief Compliance Officer',
                    'user_id': 'user_cco',
                    'authority': 'FINAL_POLICY_APPROVAL',
                    'order': 1,
                }
            ],
            'escalation_days': 5,
        },
        headers={'x-org-id': 'test-org', 'x-user-role': 'CCO'},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload['approval_chain']['approving_authorities'][0]['authority'] == 'FINAL_POLICY_APPROVAL'
    assert payload['escalation_days'] == 5
    assert payload['persisted'] is False
    app.dependency_overrides.clear()


def test_nominated_team_member_role_assignment_is_available_in_local_mode() -> None:
    app.dependency_overrides[get_db] = override_db_none
    client = TestClient(app)

    response = client.post(
        '/api/v1/organizations/current/nominated-team-members',
        json={'user_id': 'user_reviewer'},
        headers={'x-org-id': 'test-org', 'x-user-role': 'CCO'},
    )

    assert response.status_code == 201
    payload = response.json()
    assert payload == {
        'auth_org_id': 'test-org',
        'user_id': 'user_reviewer',
        'role': 'nominated_team_member',
        'assigned': False,
    }
    app.dependency_overrides.clear()


def test_storage_paths_are_org_scoped_and_supabase_urls_are_normalized() -> None:
    path = storage_path('org_example', 'policies',
                        'policy-123', 'written policy.docx')

    assert path == 'org_example/policies/policy-123/written_policy.docx'
    assert storage_url(Settings(
        SUPABASE_URL='https://project.supabase.co/rest/v1/')) == 'https://project.supabase.co'
    assert parse_storage_uri('storage://policy-documents/org_example/policies/policy-123/file.docx') == StoredFile(
        bucket='policy-documents',
        path='org_example/policies/policy-123/file.docx',
    )


def test_policy_persistence_keys_vector_chunks_by_org_and_policy() -> None:
    class Query:
        def filter(self, *_args):
            return self

        def update(self, *_args, **_kwargs):
            return 0

        def all(self):
            return []

        def delete(self, *_args, **_kwargs):
            return 0

    class RecordingSession:
        def __init__(self):
            self.added: list[object] = []
            self.flush_snapshots: list[list[object]] = []
            self.committed = False

        def query(self, *_args):
            return Query()

        def add(self, value: object):
            self.added.append(value)

        def flush(self):
            self.flush_snapshots.append(self.added.copy())

        def commit(self):
            self.committed = True

        def rollback(self):
            return None

    db = RecordingSession()
    policy_id = str(uuid4())
    embedding = [0.125] * 3072
    persist_policy_upload(
        db=db,  # type: ignore[arg-type]
        policy_id=policy_id,
        org_id='test-org',
        policy_name='Written Supervisory Procedures',
        version=1,
        file_url='storage://policy-documents/test-org/policies/policy-id/policy.docx',
        chunks=[
            TextChunk(
                index=0,
                content='Policy content',
                metadata={'filename': 'policy.docx'},
            )
        ],
        embeddings=[embedding],
        embedding_model='text-embedding-3-large',
    )

    assert len(db.flush_snapshots) == 1
    assert isinstance(db.flush_snapshots[0][0], Policy)
    chunks = [item for item in db.added if isinstance(item, DocumentChunk)]
    assert len(chunks) == 1
    assert chunks[0].org_id == 'test-org'
    assert chunks[0].policy_id == policy_id
    assert chunks[0].chunk_index == 0
    assert chunks[0].content == 'Policy content'
    assert chunks[0].document_type == 'uploaded'
    assert chunks[0].policy_name == 'Written Supervisory Procedures'
    assert chunks[0].embedding == embedding
    assert chunks[0].metadata_json == {
        'filename': 'policy.docx',
        'policy_type': 'uploaded',
        'embedding_model': 'text-embedding-3-large',
        'embedded': True,
    }
    assert db.committed is True
