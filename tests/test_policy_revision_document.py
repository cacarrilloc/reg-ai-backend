from docx import Document

from app.services.policy_revision_document import (
    add_body_paragraphs,
    normalize_body_text_for_lists,
)


def test_normalize_body_text_for_lists_splits_inline_asterisks() -> None:
    normalized = normalize_body_text_for_lists(
        "Intro text * Updated the policy * Second item"
    )

    assert normalized.splitlines() == [
        "Intro text",
        "* Updated the policy",
        "* Second item",
    ]


def test_add_body_paragraphs_uses_word_list_styles() -> None:
    document = Document()
    add_body_paragraphs(
        document,
        "Intro paragraph.\n• First bullet item\n• Second bullet item\n1. First action\n2. Second action",
    )

    styles = [paragraph.style.name for paragraph in document.paragraphs if paragraph.text.strip()]
    texts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]

    assert styles.count("List Bullet") == 2
    assert styles.count("List Number") == 2
    assert "First bullet item" in texts
    assert "First action" in texts
    assert all("•" not in text for text in texts)
    assert all("*" not in text for text in texts)


def test_add_body_paragraphs_renders_inline_asterisk_bullets() -> None:
    document = Document()
    add_body_paragraphs(
        document,
        "Policy updates required: * Updated the policy with the latest changes.",
    )

    rows = [
        (paragraph.style.name, paragraph.text)
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    assert rows == [
        ("Normal", "Policy updates required:"),
        ("List Bullet", "Updated the policy with the latest changes."),
    ]

    bullet_paragraph = next(
        paragraph for paragraph in document.paragraphs if paragraph.style.name == "List Bullet"
    )
    assert bullet_paragraph._p.pPr.numPr is not None
    assert bullet_paragraph.paragraph_format.left_indent is not None
